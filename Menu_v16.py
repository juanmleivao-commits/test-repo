#!/usr/bin/env python3
"""
CrowdStrike Falcon Unified Manager
Integra módulos operativos y reportes de auditoría.
"""

import os
import sys
import re
import time
import json
import getpass
import argparse
import unicodedata
import subprocess
from datetime import datetime, timedelta, timezone
from concurrent.futures import ThreadPoolExecutor

import requests
import pandas as pd
from openpyxl import load_workbook
from tqdm import tqdm

from falconpy import Hosts, RealTimeResponse, RealTimeResponseAdmin, IOC, FalconXSandbox

import shutil
import zipfile
import py7zr
from pathlib import Path
from typing import Any, Iterator

# =========================================================================
# --- VARIABLES GLOBALES ---
# =========================================================================

# API de VirusTotal (Módulo 7) — cargada desde variable de entorno
VIRUSTOTAL_API_KEY = os.environ.get("VIRUSTOTAL_API_KEY", "")

# Rutas Excel (Módulos 2 y 5) — sobrescribibles con variables de entorno
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE_PATH_RTR2 = os.environ.get(
    "EXCEL_PATH_RTR2",
    os.path.join(_SCRIPT_DIR, "Actualizacion.xlsx")
)
SHEET_NAME_RTR2 = "WINDOWS"

EXCEL_FILE_PATH_MOD5 = os.environ.get(
    "EXCEL_PATH_MOD5",
    os.path.join(_SCRIPT_DIR, "Despliegue.xlsx")
)
SHEET_NAME_MOD5 = "WINDOWS"

# Columnas Excel (Módulo 2)
COL_HOSTNAME = "A"
COL_RTR = "B"
COL_PUT = "C"
COL_DLP = "D"
COL_DET = "E"

# Columnas Excel y Rutas (Módulo 5)
COL_HOSTNAME_M5 = "A"
COL_RTR_M5 = "B"
COL_PUT_M5 = "C"
COL_SCRIPT_M5 = "D"
COL_DET_M5 = "E"
REMOTE_WORKDIR_M5 = r"C:\temp"

# Rutas de Sensores (Módulo 6) — sobrescribibles con variables de entorno.
# DEFAULT_DEST_SENSORS define los directorios remotos de destino en los hosts Windows
# donde se depositan los instaladores antes de ejecutarlos vía RTR.
BASE_LOCAL_SENSORS = Path(os.environ.get("BASE_SENSORS_PATH", str(Path.home() / "Downloads" / "Sensores")))
DEFAULT_DEST_SENSORS = {
    "windows": os.environ.get("DEST_SENSORS_WIN", r"C:\CrowdStrike\Sensors\Windows"),
    "macos":   os.environ.get("DEST_SENSORS_MAC", r"C:\CrowdStrike\Sensors\Mac"),
    "linux":   os.environ.get("DEST_SENSORS_LIN", r"C:\CrowdStrike\Sensors\Linux"),
}

class Colors:
    RESET = '\033[0m'
    BOLD = '\033[1m'
    RED = '\033[91m'
    GREEN = '\033[92m'
    YELLOW = '\033[93m'
    BLUE = '\033[94m'
    CYAN = '\033[96m'
    SUCCESS = GREEN
    ERROR = RED
    INFO = CYAN
    WARNING = YELLOW
    RESULT = BLUE

# =========================
# UTILIDADES GLOBALES
# =========================
def clear_screen() -> None:
    os.system('cls' if os.name == 'nt' else 'clear')

def print_banner(titulo: str) -> None:
    print("=" * 70)
    print(f"{Colors.BOLD}{Colors.CYAN}" + " " * 10 + titulo)
    print(f"{Colors.RESET}=" * 70)

def normalize(v: Any) -> str:
    return "" if v is None else str(v).strip()

def remove_accents(text: str) -> str:
    nfd = unicodedata.normalize('NFD', text)
    return ''.join(char for char in nfd if unicodedata.category(char) != 'Mn')


def normalize_tag_value(value: str) -> str:
    v = str(value).strip()
    v = remove_accents(v)
    v = v.replace(" ", "_")
    v = re.sub(r"[^A-Za-z0-9._-]", "_", v)
    v = re.sub(r"_+", "_", v).strip("_")
    return v

def print_progress(current: int, total: int, prefix: str = "Progreso", current_item: str = "") -> None:
    if total <= 0:
        return
    pct = int((current / total) * 100)
    bar_len = 24
    filled = int(bar_len * current / total)
    bar = "#" * filled + "-" * (bar_len - filled)
    extra = f" - {current_item}" if current_item else ""
    print(f"{prefix}: [{bar}] {pct}% ({current}/{total}){extra}")

# =========================================================================
# LÓGICA DE CREDENCIALES UNIFICADA
# =========================================================================
def pedir_credenciales() -> tuple[str, str]:
    """Busca las credenciales en memoria, si no están, las pide y las guarda"""
    cid = os.environ.get("FALCON_CLIENT_ID")
    csec = os.environ.get("FALCON_CLIENT_SECRET")
    
    if cid and csec:
        return cid, csec

    print(f"\n{Colors.CYAN}--- Autenticación Inicial CrowdStrike ---{Colors.RESET}")
    while True:
        cid = input(f"{Colors.CYAN}Ingrese el client ID:{Colors.RESET} ").strip()
        if not cid:
            print(f"{Colors.ERROR}✗ El client ID no puede estar vacío.{Colors.RESET}\n")
            continue

        csec = getpass.getpass(f"{Colors.CYAN}Ingrese el client secret:{Colors.RESET} ").strip()
        if not csec:
            print(f"{Colors.ERROR}✗ El client secret no puede estar vacío.{Colors.RESET}\n")
            continue

        # Guardar en memoria de la sesión
        os.environ["FALCON_CLIENT_ID"] = cid
        os.environ["FALCON_CLIENT_SECRET"] = csec
        print(f"{Colors.GREEN}Credenciales guardadas en la sesión actual.{Colors.RESET}")
        return cid, csec

def validar_credenciales_falconpy(rtr: RealTimeResponse) -> None:
    """Valida credenciales con FalconPy"""
    _ = rtr.headers
    if not getattr(rtr, "token_valid", False):
        status = getattr(rtr, "token_status", "unknown")
        reason = getattr(rtr, "token_fail_reason", "unknown")
        raise SystemExit(
            f"{Colors.ERROR}✗ Credenciales inválidas o sin permisos. "
            f"token_status={status} reason={reason}{Colors.RESET}"
        )

# =========================================================================
# MENÚ PRINCIPAL
# =========================================================================
def menu_principal() -> str:
    clear_screen()
    print(f"\n{Colors.BOLD}{Colors.CYAN}")
    print("╔" + "═" * 68 + "╗")
    print("║" + " " * 68 + "║")
    print("║" + "IBM & CROWDSTRIKE FALCON UNIFIED MANAGER".center(68) + "║")
    print("║" + " " * 68 + "║")
    print("╚" + "═" * 68 + "╝")
    print(f"{Colors.RESET}")
    
    print(f"\n{Colors.YELLOW}[ MÓDULOS OPERATIVOS ]{Colors.RESET}")
    print(f"{Colors.INFO}1.{Colors.RESET} Gestión de Tags (add/remove)               {Colors.GREEN}| File: tags.xlsx{Colors.RESET}")
    print(f"{Colors.INFO}2.{Colors.RESET} Operaciones RTR (PUT / DLP)                {Colors.GREEN}| File: Actualizacion.xlsx{Colors.RESET}")
    print(f"{Colors.INFO}3.{Colors.RESET} IOC Manager (hashes MD5/SHA256)")
    print(f"{Colors.INFO}4.{Colors.RESET} Sandbox (Win11)                            {Colors.GREEN}| Host: azsdomina01{Colors.RESET}")
    print(f"{Colors.INFO}5.{Colors.RESET} Response scripts and files                 {Colors.GREEN}| File: Despliegue.xlsx{Colors.RESET}")
    print(f"{Colors.INFO}6.{Colors.RESET} Cargar Sensores OS (Windows/Mac/Linux)")
    print(f"{Colors.INFO}7.{Colors.RESET} Validar Hash en VirusTotal")
    print(f"{Colors.INFO}8.{Colors.RESET} Delete Hash (RTR hash + confirmación)")
    print(f"{Colors.INFO}9.{Colors.RESET} Depurar Hosts Duplicados en Consola")
    
    print(f"\n{Colors.GREEN}[ REPORTES Y AUDITORÍA ]{Colors.RESET}")
    print(f"{Colors.INFO}10.{Colors.RESET} Reporte: Riesgo de Identidad             {Colors.GREEN}| Identity_Risk_v2.py{Colors.RESET}")
    print(f"{Colors.INFO}11.{Colors.RESET} Reporte: Políticas de Prevención         {Colors.GREEN}| Politicas_Actuales_v2.ps1{Colors.RESET}")
    print(f"{Colors.INFO}12.{Colors.RESET} Reporte: Exclusiones Device Control      {Colors.GREEN}| Report_CS_DeviceControl.ps1{Colors.RESET}")
    print(f"{Colors.INFO}13.{Colors.RESET} Reporte: Exclusiones Globales (ML/SV/IOA){Colors.GREEN}| Reporte_Exclusiones_Globales.ps1{Colors.RESET}")
    print(f"{Colors.INFO}14.{Colors.RESET} Reporte: Accesos de Usuarios a Consola   {Colors.GREEN}| Reporte_Usuarios_Consola.ps1{Colors.RESET}")

    print(f"\n{Colors.RED}[ SISTEMA ]{Colors.RESET}")
    print(f"{Colors.INFO}0.{Colors.RESET} Salir y limpiar credenciales")

    print(f"{Colors.BLUE}{'='*70}{Colors.RESET}")

    opcion = input(f"{Colors.BOLD}Selecciona la Opción del 0 al 14:{Colors.RESET} ").strip()
    return opcion

# =========================
# MÓDULO 1: TAGS MANAGEMENT
# =========================
REQUIRED_COLUMNS = [
    "HOSTNAME",
    "AID",
    "NOMBRE_CENTRO_DE_COSTOS",
    "EQUIPO_O_COE",
    "CELULA_COE_O_SMT",
    "EQUIPO__O_REGIÓN",
    "ECOSISTEMA_O_EQUIPO",
    "CATT",
    "TIPO_DE_EMPLEADO",
    "Relacion_Laboral",
]

STATUS_COL = "PASSO O FALLO"
DETAIL_COL = "DETALLE"


class CrowdStrikeTagManager:
    def __init__(self, client_id: str, client_secret: str, region: str = "us-1") -> None:
        self.client_id = client_id
        self.client_secret = client_secret
        self.region = region.lower()
        self.region_map = {
            "us-1": "https://api.crowdstrike.com",
            "us-2": "https://api.us-2.crowdstrike.com",
        }
        if self.region not in self.region_map:
            raise ValueError(f"Región no soportada: {region}")
        self.base_url = self.region_map[self.region]
        self.token = None

    def authenticate(self) -> bool:
        auth_url = f"{self.base_url}/oauth2/token"
        payload = {"client_id": self.client_id, "client_secret": self.client_secret}
        r = requests.post(auth_url, data=payload, timeout=20)
        r.raise_for_status()
        data = r.json()
        self.token = data.get("access_token")
        return bool(self.token)

    def get_headers(self) -> dict[str, str]:
        return {"Authorization": f"Bearer {self.token}", "Content-Type": "application/json"}

    def update_device_tags(self, device_ids: list[str], tags: list[str], action: str = "add") -> dict:
        endpoint = f"{self.base_url}/devices/entities/devices/tags/v1"
        payload = {"action": action, "device_ids": device_ids, "tags": tags}
        r = requests.patch(endpoint, json=payload, headers=self.get_headers(), timeout=60)
        if r.status_code not in (200, 202):
            raise RuntimeError(f"HTTP {r.status_code}: {r.text}")
        data = r.json()
        if data.get("errors"):
            raise RuntimeError(f"API errors: {data['errors']}")
        return data


def seleccionar_accion_tags() -> str:
    while True:
        print(f"\n{Colors.BLUE}¿Qué deseas hacer con los tags?{Colors.RESET}")
        print(f"{Colors.INFO}1.{Colors.RESET} Agregar tags (action=add)")
        print(f"{Colors.INFO}2.{Colors.RESET} Eliminar tags (action=remove)")
        op = input(f"{Colors.CYAN}Opción (1/2):{Colors.RESET} ").strip()
        if op == "1":
            return "add"
        if op == "2":
            return "remove"
        print(f"{Colors.ERROR}Opción inválida.{Colors.RESET}")


def build_tags_from_row(row: dict) -> list[str]:
    tags = []
    for col in REQUIRED_COLUMNS:
        if col in ("HOSTNAME", "AID"):
            continue
        raw = row.get(col, None)
        if raw is None:
            continue
        val = str(raw).strip()
        if not val or val.lower() == "nan":
            continue
        tag_val = normalize_tag_value(val)
        if tag_val:
            tags.append(f"FalconGroupingTags/{tag_val}")
    return list(dict.fromkeys(tags))


def run_tags_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("CrowdStrike Falcon Tags Manager (Excel)")

    parser = argparse.ArgumentParser(description="Aplica o elimina FalconGroupingTags desde Excel.")
    parser.add_argument("--excel", default="tags.xlsx", help="Ruta del Excel (default: tags.xlsx).")
    parser.add_argument("--sheet", default=0, help="Nombre hoja o índice (default 0).")
    parser.add_argument("--out", default="tags_resultado.xlsx", help="Excel de salida (default: tags_resultado.xlsx).")
    parser.add_argument("--dry-run", action="store_true", help="Dry-run: no llama a la API.")

    try:
        args = parser.parse_args()
    except SystemExit:
        args = parser.parse_args([])

    excel_path = args.excel
    if not os.path.isabs(excel_path):
        excel_path = os.path.join(os.getcwd(), excel_path)
    if not os.path.exists(excel_path):
        print(f"{Colors.ERROR}No existe el Excel: {excel_path}{Colors.RESET}")
        return

    df = pd.read_excel(excel_path, sheet_name=args.sheet, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        print(f"{Colors.ERROR}Faltan columnas: {missing}{Colors.RESET}")
        return

    if STATUS_COL not in df.columns:
        df[STATUS_COL] = ""
    if DETAIL_COL not in df.columns:
        df[DETAIL_COL] = ""

    action = seleccionar_accion_tags()
    print(f"\n{Colors.INFO}Acción seleccionada:{Colors.RESET} {action}")

    print(f"\n{Colors.BLUE}SELECCIONA TU REGIÓN:{Colors.RESET}")
    print(f"{Colors.INFO}1. US-1 (Default){Colors.RESET}")
    print(f"{Colors.INFO}2. US-2{Colors.RESET}")
    region_choice = input(f"{Colors.CYAN}Región (1-2, default 1):{Colors.RESET} ").strip()
    region = "us-2" if region_choice == "2" else "us-1"

    manager = CrowdStrikeTagManager(client_id, client_secret, region)
    try:
        print(f"\n{Colors.INFO}Autenticando...{Colors.RESET}")
        if not manager.authenticate():
            raise RuntimeError("No se pudo obtener token.")
        print(f"{Colors.SUCCESS}Autenticación OK{Colors.RESET}")
    except Exception as e:
        print(f"{Colors.ERROR}Error autenticando: {e}{Colors.RESET}")
        return

    ok = 0
    fail = 0
    skipped = 0

    for idx, row in df.iterrows():
        aid = str(row.get("AID", "") or "").strip()
        hostname = str(row.get("HOSTNAME", "") or "").strip()

        if not aid:
            df.at[idx, STATUS_COL] = "FALLO"
            df.at[idx, DETAIL_COL] = "AID vacío"
            fail += 1
            continue

        tags = build_tags_from_row(row)
        if not tags:
            df.at[idx, STATUS_COL] = "PASSO"
            df.at[idx, DETAIL_COL] = "Sin tags (todas las columnas vacías)"
            skipped += 1
            continue

        try:
            if args.dry_run:
                df.at[idx, STATUS_COL] = "PASSO"
                df.at[idx, DETAIL_COL] = f"DRY-RUN {action}: {len(tags)} tags"
            else:
                manager.update_device_tags([aid], tags, action=action)
                df.at[idx, STATUS_COL] = "PASSO"
                df.at[idx, DETAIL_COL] = f"OK {action}: {len(tags)} tags"

            ok += 1
            op = "+" if action == "add" else "-"
            print(f"{Colors.SUCCESS}PASSO{Colors.RESET} {hostname or aid} ({aid}) {op}{len(tags)} tags")
        except Exception as e:
            df.at[idx, STATUS_COL] = "FALLO"
            df.at[idx, DETAIL_COL] = str(e)[:500]
            fail += 1
            print(f"{Colors.ERROR}FALLO{Colors.RESET} {hostname or aid} ({aid}): {e}")

    out_path = args.out
    if not os.path.isabs(out_path):
        out_path = os.path.join(os.getcwd(), out_path)

    df.to_excel(out_path, index=False)
    print(f"\n{Colors.RESULT}Salida: {out_path}{Colors.RESET}")
    print(f"{Colors.RESULT}PASSO: {ok} | FALLO: {fail} | Sin tags: {skipped}{Colors.RESET}")


# =========================
# MÓDULO 2: RTR OPERATIONS
# =========================
COMMAND_CD = r"cd distribuciones"
COMMAND_PUT = r"put DeviceComplianceLocalOnboardingScript.cmd"
COMMAND_DLP = r"runscript -CloudFile=Tuya_Windows_Ejecutar_DLP_Onboarding"

INIT_TIMEOUT_SECONDS = 120
INIT_TIMEOUT_DURATION = "120s"
CMD_TIMEOUT_SECONDS = 240
CMD_TIMEOUT_DURATION = "240s"

SAVE_EVERY = 3
PUT_OK_VALUES = {"Copiado correctamente", "✅ COPIADO", "✅ Copiado", "OK", "SUCCESS"}

def ensure_headers(sheet: Any) -> None:
    if sheet[f"{COL_RTR}1"].value is None:
        sheet[f"{COL_RTR}1"] = "Estado RTR"
    if sheet[f"{COL_PUT}1"].value is None:
        sheet[f"{COL_PUT}1"] = "Estado PUT"
    if sheet[f"{COL_DLP}1"].value is None:
        sheet[f"{COL_DLP}1"] = "Estado DLP"
    if sheet[f"{COL_DET}1"].value is None:
        sheet[f"{COL_DET}1"] = "Detalles"


def should_run_dlp_based_on_put(put_status: Any) -> bool:
    return normalize(put_status) in PUT_OK_VALUES


def iter_host_rows(sheet: Any) -> Iterator[tuple[int, str]]:
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=1):
        cell = row_cells[0]
        hostname = normalize(cell.value)
        if not hostname or hostname.lower() == "hostname":
            continue
        yield cell.row, hostname


def count_hosts(sheet: Any) -> int:
    return sum(1 for _ in iter_host_rows(sheet))


def get_host_id_by_hostname(hosts: Hosts, hostname: str) -> str | None:
    resp = hosts.query_devices_by_filter_scroll(filter=f"hostname:'{hostname}'")
    resources = resp.get("body", {}).get("resources") or []
    return resources[0] if resources else None


def init_rtr_session(rtr: RealTimeResponse, host_id: str) -> dict:
    return rtr.batch_init_sessions(
        host_ids=[host_id],
        timeout=INIT_TIMEOUT_SECONDS,
        timeout_duration=INIT_TIMEOUT_DURATION
    )


def rtr_cd(rtr_admin: RealTimeResponseAdmin, batch_id: str) -> dict:
    return rtr_admin.batch_admin_command(
        base_command="cd",
        batch_id=batch_id,
        command_string=COMMAND_CD,
        timeout=CMD_TIMEOUT_SECONDS,
        timeout_duration=CMD_TIMEOUT_DURATION
    )


def rtr_put(rtr_admin: RealTimeResponseAdmin, batch_id: str) -> dict:
    return rtr_admin.batch_admin_command(
        base_command="put",
        batch_id=batch_id,
        command_string=COMMAND_PUT,
        timeout=CMD_TIMEOUT_SECONDS,
        timeout_duration=CMD_TIMEOUT_DURATION
    )


def rtr_runscript_dlp(rtr_admin: RealTimeResponseAdmin, batch_id: str) -> dict:
    return rtr_admin.batch_admin_command(
        base_command="runscript",
        batch_id=batch_id,
        command_string=COMMAND_DLP,
        timeout=CMD_TIMEOUT_SECONDS,
        timeout_duration=CMD_TIMEOUT_DURATION
    )


def run_put_for_all(sheet: Any, workbook: Any, hosts: Hosts, rtr: RealTimeResponse, rtr_admin: RealTimeResponseAdmin) -> None:
    total = count_hosts(sheet)
    done = 0
    saves = 0

    for row, hostname in iter_host_rows(sheet):
        done += 1
        print_progress(done, total, prefix="PUT", current_item=hostname)

        try:
            sheet[f"{COL_RTR}{row}"] = ""
            sheet[f"{COL_PUT}{row}"] = ""
            sheet[f"{COL_DET}{row}"] = ""

            host_id = get_host_id_by_hostname(hosts, hostname)
            if not host_id:
                sheet[f"{COL_RTR}{row}"] = "❌ No encontrado"
                sheet[f"{COL_PUT}{row}"] = "N/A"
                sheet[f"{COL_DET}{row}"] = "Hostname no existe en CrowdStrike"
                continue

            session_init = init_rtr_session(rtr, host_id)
            sc = session_init.get("status_code")

            if sc == 404:
                sheet[f"{COL_RTR}{row}"] = "❌ OFFLINE"
                sheet[f"{COL_PUT}{row}"] = "⏭️ Omitido"
                sheet[f"{COL_DET}{row}"] = "Host offline"
                continue

            if sc not in (200, 201):
                sheet[f"{COL_RTR}{row}"] = f"❌ Error RTR {sc}"
                sheet[f"{COL_PUT}{row}"] = "❌ No ejecutado"
                sheet[f"{COL_DET}{row}"] = "Falló batch_init_sessions"
                continue

            sheet[f"{COL_RTR}{row}"] = "✅ Conectado"
            batch_id = session_init["body"]["batch_id"]

            cd_resp = rtr_cd(rtr_admin, batch_id)
            if cd_resp.get("status_code") not in (200, 201):
                sheet[f"{COL_PUT}{row}"] = f"❌ Error CD {cd_resp.get('status_code')}"
                sheet[f"{COL_DET}{row}"] = "Falló cambio de directorio"
                continue

            put_resp = rtr_put(rtr_admin, batch_id)
            if put_resp.get("status_code") in (200, 201):
                sheet[f"{COL_PUT}{row}"] = "Copiado correctamente"
                sheet[f"{COL_DET}{row}"] = "PUT OK"
            else:
                sheet[f"{COL_PUT}{row}"] = "Falló copiado"
                sheet[f"{COL_DET}{row}"] = f"PUT error {put_resp.get('status_code')}"

        except KeyboardInterrupt:
            workbook.save(EXCEL_FILE_PATH_RTR2)
            raise
        except Exception as ex:
            sheet[f"{COL_RTR}{row}"] = "❌ Excepción"
            sheet[f"{COL_PUT}{row}"] = "Falló copiado"
            sheet[f"{COL_DET}{row}"] = f"Excepción: {str(ex)[:160]}"

        if done % SAVE_EVERY == 0:
            workbook.save(EXCEL_FILE_PATH_RTR2)
            saves += 1
            print(f"{Colors.INFO}Guardado parcial Excel (#{saves}){Colors.RESET}")

        time.sleep(1)

    workbook.save(EXCEL_FILE_PATH_RTR2)
    print(f"{Colors.SUCCESS}PUT finalizado y Excel guardado.{Colors.RESET}")


def run_dlp_only_successful_put(sheet: Any, workbook: Any, hosts: Hosts, rtr: RealTimeResponse, rtr_admin: RealTimeResponseAdmin) -> None:
    candidates = [r for r, _h in iter_host_rows(sheet)
                  if should_run_dlp_based_on_put(sheet[f"{COL_PUT}{r}"].value)]
    total = len(candidates)
    done = 0
    saves = 0

    for row in candidates:
        hostname = normalize(sheet[f"{COL_HOSTNAME}{row}"].value)
        done += 1
        print_progress(done, total, prefix="DLP", current_item=hostname)

        try:
            sheet[f"{COL_DLP}{row}"] = ""

            host_id = get_host_id_by_hostname(hosts, hostname)
            if not host_id:
                sheet[f"{COL_DLP}{row}"] = "❌ No encontrado"
                sheet[f"{COL_DET}{row}"] = "Hostname no existe en CrowdStrike"
                continue

            session_init = init_rtr_session(rtr, host_id)
            sc = session_init.get("status_code")

            if sc == 404:
                sheet[f"{COL_DLP}{row}"] = "⏭️ OFFLINE"
                sheet[f"{COL_DET}{row}"] = "Host offline"
                continue

            if sc not in (200, 201):
                sheet[f"{COL_DLP}{row}"] = f"❌ Error RTR {sc}"
                sheet[f"{COL_DET}{row}"] = "Falló batch_init_sessions"
                continue

            batch_id = session_init["body"]["batch_id"]

            cd_resp = rtr_cd(rtr_admin, batch_id)
            if cd_resp.get("status_code") not in (200, 201):
                sheet[f"{COL_DLP}{row}"] = f"❌ Error CD {cd_resp.get('status_code')}"
                sheet[f"{COL_DET}{row}"] = "Falló cambio de directorio"
                continue

            dlp_resp = rtr_runscript_dlp(rtr_admin, batch_id)
            if dlp_resp.get("status_code") in (200, 201):
                sheet[f"{COL_DLP}{row}"] = "✅ Ejecutado"
                sheet[f"{COL_DET}{row}"] = "DLP runscript enviado OK"
            else:
                sheet[f"{COL_DLP}{row}"] = f"❌ Error {dlp_resp.get('status_code')}"
                sheet[f"{COL_DET}{row}"] = "Falló runscript DLP"

        except KeyboardInterrupt:
            workbook.save(EXCEL_FILE_PATH_RTR2)
            raise
        except Exception as ex:
            sheet[f"{COL_DLP}{row}"] = "❌ Excepción"
            sheet[f"{COL_DET}{row}"] = f"Excepción: {str(ex)[:160]}"

        if done % SAVE_EVERY == 0:
            workbook.save(EXCEL_FILE_PATH_RTR2)
            saves += 1
            print(f"{Colors.INFO}Guardado parcial Excel (#{saves}){Colors.RESET}")

        time.sleep(1)

    workbook.save(EXCEL_FILE_PATH_RTR2)
    print(f"{Colors.SUCCESS}DLP finalizado y Excel guardado.{Colors.RESET}")


def menu_rtr_operations() -> str:
    while True:
        print(f"\n{Colors.BLUE}OPERACIONES RTR{Colors.RESET}")
        print(f"{Colors.INFO}1.{Colors.RESET} PUT (copiar archivo) a todos")
        print(f"{Colors.INFO}2.{Colors.RESET} Ejecutar DLP SOLO en PUT OK")
        print(f"{Colors.INFO}3.{Colors.RESET} Ambos (PUT y luego DLP en OK)")
        print(f"{Colors.INFO}4.{Colors.RESET} Volver")
        opcion = input(f"{Colors.CYAN}Opción (1-4):{Colors.RESET} ").strip()
        return opcion


def run_rtr_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("CrowdStrike Falcon RTR Operations")

    rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
    validar_credenciales_falconpy(rtr)
    print(f"{Colors.SUCCESS}✓ Autenticación exitosa{Colors.RESET}")

    rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)
    hosts = Hosts(auth_object=rtr.auth_object)

    while True:
        opcion = menu_rtr_operations()

        if opcion == "4":
            print(f"{Colors.INFO}Volviendo al menú principal...{Colors.RESET}")
            break

        workbook = load_workbook(filename=EXCEL_FILE_PATH_RTR2)
        sheet = workbook[SHEET_NAME_RTR2]
        ensure_headers(sheet)
        workbook.save(EXCEL_FILE_PATH_RTR2)

        try:
            if opcion == "1":
                run_put_for_all(sheet, workbook, hosts, rtr, rtr_admin)
            elif opcion == "2":
                run_dlp_only_successful_put(sheet, workbook, hosts, rtr, rtr_admin)
            elif opcion == "3":
                run_put_for_all(sheet, workbook, hosts, rtr, rtr_admin)
                run_dlp_only_successful_put(sheet, workbook, hosts, rtr, rtr_admin)
            else:
                print(f"{Colors.ERROR}Opción inválida.{Colors.RESET}")
        finally:
            workbook.save(EXCEL_FILE_PATH_RTR2)
            workbook.close()

        if opcion in ["1", "2", "3"]:
            input(f"\n{Colors.CYAN}Listo. Presione Enter para volver al menú...{Colors.RESET}")
            clear_screen()
            print_banner("CrowdStrike Falcon RTR Operations")


# =========================
# MÓDULO 3: IOC MANAGER
# =========================
HASH_REGEX_IOC = re.compile(r"^([A-Fa-f0-9]{32}|[A-Fa-f0-9]{64})$")


def is_valid_hash_ioc(hash_str: str) -> bool:
    return bool(HASH_REGEX_IOC.match(hash_str.strip()))


def get_hash_type_ioc(hash_str: str) -> str | None:
    h = hash_str.strip()
    if len(h) == 32:
        return "md5"
    if len(h) == 64:
        return "sha256"
    return None


def extract_hashes_from_excel_ioc(excel_path: str, sheet_name: str = "Reporte") -> list[str]:
    try:
        df = pd.read_excel(excel_path, sheet_name=sheet_name, dtype=str)
        hashes = []

        for col in df.columns:
            col_upper = str(col).strip().upper()
            if "VALUE" in col_upper or "HASH" in col_upper:
                for hash_val in df[col]:
                    if pd.isna(hash_val):
                        continue
                    hash_str = str(hash_val).strip()
                    if is_valid_hash_ioc(hash_str):
                        hashes.append(hash_str.lower())

        seen = set()
        uniq = []
        for h in hashes:
            if h not in seen:
                uniq.append(h)
                seen.add(h)
        return uniq

    except Exception as e:
        print(f"{Colors.ERROR}Error leyendo Excel: {e}{Colors.RESET}")
        return []


def build_ioc_payloads(hashes: list[str], requisition_number: str) -> tuple[list[dict], str, str]:
    fecha_actual = datetime.now().strftime("%d-%m-%Y")
    description = f"{fecha_actual} - Se agregan hashes correspondientes al caso {requisition_number}"

    exp_local_date = (datetime.now() + timedelta(days=365)).date()
    expiration_iso_utc = datetime(
        exp_local_date.year, exp_local_date.month, exp_local_date.day,
        0, 0, 0, tzinfo=timezone.utc
    ).strftime("%Y-%m-%dT%H:%M:%S.000Z")

    indicators = []
    for h in hashes:
        htype = get_hash_type_ioc(h)
        if not htype:
            continue

        indicators.append({
            "type": htype,
            "value": h,
            "platforms": ["windows", "mac", "linux"],
            "applied_globally": True,
            "action": "prevent",
            "severity": "informational",
            "expiration": expiration_iso_utc,
            "description": description,
            "source": "Tuya Security Team"
        })

    return indicators, description, expiration_iso_utc


def pretty(obj: Any) -> str:
    try:
        return json.dumps(obj, indent=2, ensure_ascii=False)
    except Exception:
        return str(obj)


def create_iocs_falconpy(ioc_client: IOC, indicators: list[dict], ignore_warnings: bool = True) -> tuple[list[str], list[str]]:
    created = []
    failed = []

    for idx, ioc in enumerate(indicators, 1):
        print(f"\n{Colors.CYAN}[{idx}/{len(indicators)}] {ioc['type'].upper()}: {ioc['value']}{Colors.RESET}")

        try:
            body = {"indicators": [ioc]}
            resp = ioc_client.indicator_create(body=body, ignore_warnings=ignore_warnings)
            sc = resp.get("status_code", 0)

            if sc in (200, 201, 202):
                created.append(ioc["value"])
                print(f"{Colors.SUCCESS}✓ Creado (HTTP {sc}){Colors.RESET}")
            else:
                failed.append(ioc["value"])
                print(f"{Colors.ERROR}✗ Falló (HTTP {sc}){Colors.RESET}")
                print(f"{Colors.WARNING}Respuesta completa:{Colors.RESET}")
                print(pretty(resp.get("body", resp)))

        except Exception as e:
            failed.append(ioc["value"])
            print(f"{Colors.ERROR}✗ Error: {str(e)}{Colors.RESET}")

        time.sleep(0.25)

    return created, failed


def run_ioc_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("CrowdStrike IOC Manager - PRODUCCIÓN (FINAL)")

    excel_filename = input(f"\n{Colors.CYAN}Excel file (ej: hashes.xlsx):{Colors.RESET} ").strip()
    excel_path = os.path.join(os.getcwd(), excel_filename)

    if not os.path.exists(excel_path):
        print(f"{Colors.ERROR}Archivo no encontrado: {excel_path}{Colors.RESET}")
        return

    print(f"\n{Colors.INFO}Extrayendo hashes...{Colors.RESET}")
    hashes = extract_hashes_from_excel_ioc(excel_path, sheet_name="Reporte")

    if not hashes:
        print(f"{Colors.ERROR}No se encontraron hashes válidos.{Colors.RESET}")
        return

    print(f"{Colors.SUCCESS}✓ {len(hashes)} hashes encontrados{Colors.RESET}")
    for h in hashes:
        print(f"  - {get_hash_type_ioc(h).upper()}: {h}")

    try:
        print(f"\n{Colors.INFO}Autenticando...{Colors.RESET}")
        ioc_client = IOC(client_id=client_id, client_secret=client_secret)
        _ = ioc_client.indicator_search()
        print(f"{Colors.SUCCESS}✓ Autenticación exitosa{Colors.RESET}")
    except Exception as e:
        print(f"{Colors.ERROR}✗ Error autenticando: {e}{Colors.RESET}")
        return

    ritm = input(f"\n{Colors.CYAN}Número de requisición (ej: RITM0136122):{Colors.RESET} ").strip()
    if not ritm:
        print(f"{Colors.ERROR}RITM vacío. Abortando.{Colors.RESET}")
        return

    indicators, desc, exp = build_ioc_payloads(hashes, ritm)

    print(f"\n{Colors.INFO}Se aplicarán estos metadatos:{Colors.RESET}")
    print(f"  - Action: prevent (block)")
    print(f"  - Severity: informational")
    print(f"  - Expiration: {exp}")
    print(f"  - Description: {desc}")
    print(f"  - ignore_warnings: true")

    confirmar = input(f"\n{Colors.CYAN}¿Crear {len(indicators)} IOCs? (s/n):{Colors.RESET} ").strip().lower()
    if confirmar != "s":
        print(f"{Colors.WARNING}Cancelado.{Colors.RESET}")
        return

    created, failed = create_iocs_falconpy(ioc_client, indicators, ignore_warnings=True)

    print(f"\n{Colors.RESULT}{'='*70}{Colors.RESET}")
    print(f"{Colors.SUCCESS}✓ Creados: {len(created)}/{len(indicators)}{Colors.RESET}")
    if failed:
        print(f"{Colors.ERROR}✗ Fallidos: {len(failed)}{Colors.RESET}")
        for h in failed[:50]:
            print(f"  - {h}")
    print(f"{Colors.RESULT}{'='*70}{Colors.RESET}")


# =========================
# MÓDULO 4: RTR -> ZIP -> SANDBOX (Win11) + CLEANUP
# =========================

REMOTE_DIR_SANDBOX = r"C:\proc_ent"      # FIJO, no se toca, no se borra
REMOTE_TMP_SANDBOX = r"C:\Windows\Temp"
LOCAL_WORKDIR_SANDBOX = Path("./rtr_sandbox_work")
DEFAULT_ENVIRONMENT_ID_SANDBOX = 140   # Windows 11 64-bit (140) en Falcon Sandbox
EXTENSIONS_SANDBOX = {
    ".doc", ".docx", ".ppt", ".pps", ".pptx", ".ppsx", ".xls", ".xlsx",
    ".rtf", ".pub", ".pdf", ".csv", ".exe", ".scr", ".pif", ".dll", ".com",
    ".cpl", ".ps1", ".psd1", ".psm1", ".py", ".pl", ".js", ".vbs", ".vbe",
    ".sct", ".wsf", ".hta", ".zip", ".7z", ".apk", ".jar", ".eml", ".msg",
    ".svg", ".swf", ".lnk", ".chm"
}
RTR_PASSWORD_SANDBOX = "infected"
POLL_SLEEP_SANDBOX = 2
POLL_MAX_SECONDS_SANDBOX = 900

def require_ok_simple(resp: dict, context: str) -> dict:
    sc = resp.get("status_code", 0)
    body = resp.get("body") or {}
    if isinstance(body, dict) and body.get("errors"):
        raise RuntimeError(f"{context} - API errors:\n{json.dumps(body['errors'], indent=2, ensure_ascii=False)}")
    if sc not in [200, 201, 202]:
        raise RuntimeError(f"{context} - HTTP {sc}:\n{json.dumps(resp, indent=2, ensure_ascii=False)}")
    return resp

def wait_admin_command_simple(rtr_admin: RealTimeResponseAdmin, cloud_request_id: str, max_seconds: int = POLL_MAX_SECONDS_SANDBOX) -> dict:
    start = time.time()
    while True:
        last = rtr_admin.check_admin_command_status(cloud_request_id=cloud_request_id, sequence_id=0)
        sc = last.get("status_code", 0)
        if sc not in [200, 201, 202]:
            raise RuntimeError(f"check_admin_command_status HTTP {sc}:\n{json.dumps(last, indent=2, ensure_ascii=False)}")

        resources = (last.get("body") or {}).get("resources") or []
        if resources and resources[0].get("complete") is True:
            if resources[0].get("stderr"):
                raise RuntimeError(
                    f"RTR Admin command completed with stderr:\n{resources[0].get('stderr')}\n"
                    f"Full:\n{json.dumps(last, indent=2, ensure_ascii=False)}"
                )
            return last

        if time.time() - start > max_seconds:
            raise TimeoutError(f"Timeout esperando cloud_request_id={cloud_request_id}.\n{json.dumps(last, indent=2, ensure_ascii=False)}")
        time.sleep(POLL_SLEEP_SANDBOX)

def rtr_admin_cmd_simple(rtr_admin: RealTimeResponseAdmin, session_id: str, base_command: str, command_string: str) -> str:
    # Ejecuta un comando admin RTR y retorna el cloud_request_id para polling posterior.
    resp = rtr_admin.execute_admin_command(
        base_command=base_command,
        session_id=session_id,
        command_string=command_string
    )
    require_ok_simple(resp, f"execute_admin_command {base_command}")
    resources = (resp.get("body") or {}).get("resources") or []
    # Guard ANTES de acceder a resources[0]: evita IndexError silencioso cuando la API
    # devuelve lista vacía (p.ej. sesión RTR expirada o error intermitente de la plataforma).
    if not resources or "cloud_request_id" not in resources[0]:
        raise RuntimeError(f"execute_admin_command no devolvió cloud_request_id:\n{json.dumps(resp, indent=2, ensure_ascii=False)}")
    return resources[0]["cloud_request_id"]


def _sanitize_ps_literal(value: str) -> str:
    """Sanitiza un valor para interpolación segura en un string PowerShell de comilla simple.

    Mitiga A1/A03 Injection: un nombre de archivo con comilla simple rompe el contexto
    del string en PowerShell y permite inyectar sentencias arbitrarias con privilegios RTR.
    - Rechaza traversal de ruta ('..', '/') antes de cualquier interpolación.
    - Escapa comillas simples según la convención PowerShell: '' (dos comillas simples).
    """
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Componente de ruta vacío o inválido: {value!r}")
    # Rechazar traversal antes de cualquier escape para no ocultar rutas maliciosas
    if ".." in value or "/" in value:
        raise ValueError(f"Traversal de ruta detectado en componente: {value!r}")
    # En PowerShell, la única secuencia que rompe un single-quoted string es la comilla simple;
    # el estándar PowerShell la representa duplicándola ('').
    return value.replace("'", "''")

def sanitize_filename_simple(name: str) -> str:
    from urllib.parse import unquote
    n = unquote(name)
    n = n.replace("/", "_").replace("\\", "_").strip()
    return n

def find_extracted_sha256_simple(rtr: RealTimeResponse, session_id: str, cloud_request_id: str, max_seconds: int = 300) -> tuple[str, dict]:
    start = time.time()
    last_files = None
    while True:
        files = rtr.list_files_v2(session_id=session_id)
        last_files = files
        require_ok_simple(files, "list_files_v2")
        resources = (files.get("body") or {}).get("resources") or []
        match = [f for f in resources if f.get("cloud_request_id") == cloud_request_id]
        if match:
            sha256 = match[0].get("sha256")
            if sha256:
                return sha256, match[0]

        if time.time() - start > max_seconds:
            raise RuntimeError(f"No apareció el archivo extraído.\nÚltimo list_files_v2:\n{json.dumps(last_files, indent=2, ensure_ascii=False)}")
        time.sleep(2)

def download_extracted_7z_simple(
    rtr: RealTimeResponse,
    session_id: str,
    cloud_request_id: str,
    filename: str,
    outpath: Path,
    tries: int = 12,
    sleeps: int = 5,
) -> Path:
    sha256, file_obj = find_extracted_sha256_simple(rtr, session_id, cloud_request_id, max_seconds=300)
    actual_name = file_obj.get("name")
    print(f"{Colors.INFO}[DEBUG]{Colors.RESET} Extracted file found: sha256={sha256} name={actual_name}")

    outpath.parent.mkdir(parents=True, exist_ok=True)
    last_err = None
    for attempt in range(1, tries + 1):
        try:
            with rtr.get_extracted_file_contents(
                session_id=session_id,
                sha256=sha256,
                filename=actual_name,
                stream=True
            ) as request:
                request.raise_for_status()
                with open(outpath, 'wb') as f:
                    for chunk in request.iter_content(chunk_size=1024 * 1024):
                        if chunk:
                            f.write(chunk)
            return outpath
        except Exception as e:
            last_err = str(e)
            print(f"    {Colors.WARNING}[WARN]{Colors.RESET} Intento {attempt}/{tries} fallido, esperando {sleeps}s...")
            time.sleep(sleeps)

    raise RuntimeError(f"No se pudo descargar extracted file tras {tries} intentos. Último error: {last_err}")

def extract_sha256_from_upload_simple(resp: dict) -> str:
    require_ok_simple(resp, "FalconX upload_sample")
    resources = (resp.get("body") or {}).get("resources") or []
    if not resources:
        raise RuntimeError(f"upload_sample sin resources:\n{json.dumps(resp, indent=2, ensure_ascii=False)}")
    sha256 = resources[0].get("sha256") or resources[0].get("id") or resources[0]
    if not sha256 or not isinstance(sha256, str):
        raise RuntimeError(f"No pude determinar sha256/id:\n{json.dumps(resp, indent=2, ensure_ascii=False)}")
    return sha256

def cleanup_remote_zip_simple(rtr_admin: RealTimeResponseAdmin, session_id: str, remote_zip: str) -> None:
    # Escapar comillas simples en la ruta completa antes de interpolar en PowerShell (A1 Injection).
    # remote_zip es una ruta interna construida con timestamp, pero se sanitiza por defensa en profundidad.
    safe_zip = remote_zip.replace("'", "''")
    cleanup_ps = (
        "$ErrorActionPreference='SilentlyContinue'; "
        f"if (Test-Path -LiteralPath '{safe_zip}') {{ Remove-Item -LiteralPath '{safe_zip}' -Force }}; "
        "Write-Output 'CLEANUP_OK'"
    )
    cleanup_cmd = f"runscript -Raw=```{cleanup_ps}```"
    crid_cleanup = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", cleanup_cmd)
    wait_admin_command_simple(rtr_admin, crid_cleanup)
    print(f"{Colors.SUCCESS}[OK]{Colors.RESET} Cleanup remoto: eliminado {remote_zip}")

def cleanup_local_artifacts_simple(local_7z: Path, extracted_dir: Path, unzip_dir: Path) -> None:
    for d in (unzip_dir, extracted_dir):
        if d and d.exists():
            shutil.rmtree(d, ignore_errors=True)
    if local_7z and local_7z.exists():
        try:
            local_7z.unlink()
        except Exception:
            pass

def run_rtr_zip_sandbox_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("RTR -> ZIP -> Sandbox (Windows 11) + cleanup")

    aid = input(f"{Colors.CYAN}AID del host (Host ID):{Colors.RESET} ").strip()
    if not aid:
        print(f"{Colors.ERROR}AID vacío. Abortando.{Colors.RESET}")
        return

    rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
    validar_credenciales_falconpy(rtr)
    rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)
    falconx = FalconXSandbox(client_id=client_id, client_secret=client_secret)

    LOCAL_WORKDIR_SANDBOX.mkdir(parents=True, exist_ok=True)

    session_id = None
    remote_zip = None
    local_7z = None
    extracted_dir = None
    unzip_dir = None

    try:
        sess = rtr.init_session(device_id=aid)
        require_ok_simple(sess, "init_session")
        session_id = sess["body"]["resources"][0]["session_id"]
        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} RTR session_id: {session_id}")

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        remote_zip = rf"{REMOTE_TMP_SANDBOX}\proc_ent_{stamp}.zip"
        zip_name = remote_zip.split("\\")[-1]

        ps = (
            "$ErrorActionPreference='Stop'; "
            f"if (!(Test-Path -LiteralPath '{REMOTE_DIR_SANDBOX}')) {{ throw 'No existe {REMOTE_DIR_SANDBOX}' }}; "
            f"Compress-Archive -Path '{REMOTE_DIR_SANDBOX}\\*' -DestinationPath '{remote_zip}' -Force; "
            f"Write-Output 'ZIP_OK {remote_zip}'"
        )
        cmd_runscript = f"runscript -Raw=```{ps}```"
        crid_zip = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", cmd_runscript)
        wait_admin_command_simple(rtr_admin, crid_zip)
        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} ZIP creado en host: {remote_zip}")

        crid_get = rtr_admin_cmd_simple(rtr_admin, session_id, "get", f"get '{remote_zip}'")
        wait_admin_command_simple(rtr_admin, crid_get)
        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} ZIP enviado a CrowdStrike cloud (get)")

        local_7z = LOCAL_WORKDIR_SANDBOX / f"{zip_name}.7z"
        download_extracted_7z_simple(rtr, session_id, crid_get, zip_name, local_7z)
        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} Descargado: {local_7z}")

        extracted_dir = LOCAL_WORKDIR_SANDBOX / f"extracted_{stamp}"
        extracted_dir.mkdir(parents=True, exist_ok=True)
        with py7zr.SevenZipFile(local_7z, mode='r', password=RTR_PASSWORD_SANDBOX) as z:
            z.extractall(path=extracted_dir)

        zips = list(extracted_dir.rglob("*.zip"))
        if not zips:
            raise RuntimeError(f"No encontré .zip dentro de {local_7z}")

        local_zip = zips[0]
        unzip_dir = LOCAL_WORKDIR_SANDBOX / f"unzipped_{stamp}"
        unzip_dir.mkdir(parents=True, exist_ok=True)

        with zipfile.ZipFile(local_zip, 'r') as zf:
            zf.extractall(unzip_dir)

        changed = True
        while changed:
            changed = False
            for p in list(unzip_dir.rglob("*")):
                if not p.is_file():
                    continue
                suffix = p.suffix.lower()
                target_dir = p.parent / p.stem
                try:
                    if suffix == ".zip":
                        target_dir.mkdir(parents=True, exist_ok=True)
                        with zipfile.ZipFile(p, "r") as zf:
                            zf.extractall(target_dir)
                        p.unlink() 
                        changed = True
                    elif suffix == ".7z":
                        target_dir.mkdir(parents=True, exist_ok=True)
                        with py7zr.SevenZipFile(p, mode="r", password=RTR_PASSWORD_SANDBOX) as zz:
                            zz.extractall(path=target_dir)
                        p.unlink()
                        changed = True
                except Exception as e:
                    print(f"{Colors.WARNING}[WARN]{Colors.RESET} No se pudo extraer {p.name}: {e}")

        FINAL_EXTENSIONS = {e for e in EXTENSIONS_SANDBOX if e not in ('.zip', '.7z')}
        candidates = [p for p in unzip_dir.rglob("*") if p.is_file() and p.suffix.lower() in FINAL_EXTENSIONS]

        if not candidates:
            print(f"{Colors.WARNING}[INFO]{Colors.RESET} No se encontraron archivos válidos para sandbox en la ruta de extracción.")
            return

        from urllib.parse import unquote
        print(f"{Colors.INFO}[INFO]{Colors.RESET} Archivos a enviar al sandbox: {len(candidates)}")
        for p in candidates:
            print(f"  - {unquote(p.name)}")

        for p in candidates:
            safe_name = sanitize_filename_simple(p.name)
            up = falconx.upload_sample(
                sample=p.read_bytes(),
                file_name=safe_name,
                comment=f"Auto-upload from RTR folder {REMOTE_DIR_SANDBOX}, is_confidential=True"
            )
            sha256 = extract_sha256_from_upload_simple(up)

            sub = falconx.submit(
                sha256=sha256,
                submit_name=safe_name,
                aid=aid,
                environment_id=DEFAULT_ENVIRONMENT_ID_SANDBOX
            )
            sc = sub.get("status_code", 0)
            if sc not in [200, 201, 202]:
                print(f"{Colors.ERROR}[WARN]{Colors.RESET} Submit falló para {safe_name} (HTTP {sc})")
                print(json.dumps(sub, indent=2, ensure_ascii=False))
            else:
                print(f"{Colors.SUCCESS}[OK]{Colors.RESET} Submitted: {safe_name} (env_id={DEFAULT_ENVIRONMENT_ID_SANDBOX})")

    finally:
        if session_id and remote_zip:
            try:
                cleanup_remote_zip_simple(rtr_admin, session_id, remote_zip)
            except Exception as e:
                print(f"{Colors.WARNING}[WARN]{Colors.RESET} Cleanup remoto falló: {e}")

        try:
            cleanup_local_artifacts_simple(local_7z, extracted_dir, unzip_dir)
        except Exception as e:
            print(f"{Colors.WARNING}[WARN]{Colors.RESET} Cleanup local falló: {e}")

        if session_id:
            try:
                rtr.delete_session(session_id=session_id)
            except Exception:
                pass

        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} Finalizado.")


# =========================
# MÓDULO 5: RTR MASIVO SIMPLIFICADO (Excel)
# =========================
def ensure_headers_mod5(sheet: Any) -> None:
    if sheet[f"{COL_RTR_M5}1"].value is None: sheet[f"{COL_RTR_M5}1"] = "Estado RTR"
    if sheet[f"{COL_PUT_M5}1"].value is None: sheet[f"{COL_PUT_M5}1"] = "Estado PUT"
    if sheet[f"{COL_SCRIPT_M5}1"].value is None: sheet[f"{COL_SCRIPT_M5}1"] = "Estado Script"
    if sheet[f"{COL_DET_M5}1"].value is None: sheet[f"{COL_DET_M5}1"] = "Detalles"


def iter_host_rows_mod5(sheet: Any) -> Iterator[tuple[int, str]]:
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=1):
        cell = row_cells[0]
        hostname = normalize(cell.value)
        if not hostname or hostname.lower() == "hostname":
            continue
        yield cell.row, hostname


def count_hosts_mod5(sheet: Any) -> int:
    return sum(1 for _ in iter_host_rows_mod5(sheet))


def get_host_id_mod5(hosts: Hosts, hostname: str) -> str | None:
    resp = hosts.query_devices_by_filter_scroll(filter=f"hostname:'{hostname}'")
    resources = resp.get("body", {}).get("resources") or []
    return resources[0] if resources else None


def rtr_mkdir_mod5(rtr_admin: RealTimeResponseAdmin, batch_id: str) -> dict:
    ps_cmd = f"New-Item -Path '{REMOTE_WORKDIR_M5}' -ItemType Directory -Force"
    cmd = f"runscript -Raw=```{ps_cmd}```"
    return rtr_admin.batch_admin_command(
        base_command="runscript",
        batch_id=batch_id,
        command_string=cmd,
        timeout=CMD_TIMEOUT_SECONDS,
        timeout_duration=CMD_TIMEOUT_DURATION
    )


def run_rtr_universal_module(client_id: str, client_secret: str) -> None:
    from openpyxl import load_workbook
    clear_screen()
    print_banner("Despliegue Masivo (Excel)")

    print(f"{Colors.INFO}Deja en blanco si no quieres usar esa opción.{Colors.RESET}")
    put_file = input(f"{Colors.CYAN}Nombre del archivo PUT (ej: agent.msi):{Colors.RESET} ").strip()
    
    script_name = ""
    script_type = ""
    
    print(f"\n{Colors.BLUE}Tipo de Script:{Colors.RESET}")
    print("1. Custom Script (CloudFile) [Default]")
    print("2. Falcon Script")
    t_op = input("Opción (1/2): ").strip()
    
    script_type = "falcon" if t_op == "2" else "custom"
    label = "CloudFile" if script_type == "custom" else "FalconScript Name"
    
    script_name = input(f"{Colors.CYAN}Nombre del Script ({label}):{Colors.RESET} ").strip()

    if not put_file and not script_name:
        print(f"\n{Colors.ERROR}No ingresaste nada. Abortando.{Colors.RESET}")
        time.sleep(2)
        return

    print(f"\n{Colors.YELLOW}--- RESUMEN ---{Colors.RESET}")
    print(f"Excel:   {EXCEL_FILE_PATH_MOD5}")
    print(f"PUT:     {put_file if put_file else '(Ninguno)'}")
    print(f"Script:  {script_name if script_name else '(Ninguno)'} [{script_type}]")
    print(f"Ruta:    {REMOTE_WORKDIR_M5}")
    
    if input(f"\n{Colors.CYAN}¿Iniciar despliegue? (s/n):{Colors.RESET} ").lower() != 's':
        return

    rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
    validar_credenciales_falconpy(rtr)
    rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)
    hosts = Hosts(auth_object=rtr.auth_object)

    try:
        workbook = load_workbook(filename=EXCEL_FILE_PATH_MOD5)
        sheet = workbook[SHEET_NAME_MOD5]
        ensure_headers_mod5(sheet)
    except Exception as e:
        print(f"{Colors.ERROR}Error abriendo Excel: {e}{Colors.RESET}")
        return

    total = count_hosts_mod5(sheet)
    done = 0
    saves = 0

    print(f"\n{Colors.INFO}Procesando {total} hosts...{Colors.RESET}")

    for row, hostname in iter_host_rows_mod5(sheet):
        done += 1
        print_progress(done, total, prefix="RTR", current_item=hostname)
        
        sheet[f"{COL_RTR_M5}{row}"] = ""
        sheet[f"{COL_PUT_M5}{row}"] = ""
        sheet[f"{COL_SCRIPT_M5}{row}"] = ""
        sheet[f"{COL_DET_M5}{row}"] = ""

        try:
            aid = get_host_id_mod5(hosts, hostname)
            if not aid:
                sheet[f"{COL_RTR_M5}{row}"] = "❌ No encontrado"
                continue

            sess = rtr.batch_init_sessions(
                host_ids=[aid], 
                timeout=INIT_TIMEOUT_SECONDS, 
                timeout_duration=INIT_TIMEOUT_DURATION
            )
            
            if sess.get("status_code") == 404:
                sheet[f"{COL_RTR_M5}{row}"] = "❌ OFFLINE"
                continue
            
            if sess.get("status_code") not in (200, 201):
                sheet[f"{COL_RTR_M5}{row}"] = f"Error {sess.get('status_code')}"
                continue

            batch_id = sess["body"]["batch_id"]
            sheet[f"{COL_RTR_M5}{row}"] = "✅ Conectado"

            if put_file:
                rtr_mkdir_mod5(rtr_admin, batch_id) 
                
                rtr_admin.batch_admin_command(
                    base_command="cd", batch_id=batch_id, command_string=f"cd {REMOTE_WORKDIR_M5}"
                )

                put_resp = rtr_admin.batch_admin_command(
                    base_command="put",
                    batch_id=batch_id,
                    command_string=f"put {put_file}",
                    timeout=CMD_TIMEOUT_SECONDS,
                    timeout_duration=CMD_TIMEOUT_DURATION
                )

                if put_resp.get("status_code") in (200, 201):
                    sheet[f"{COL_PUT_M5}{row}"] = "✅ Copiado"
                else:
                    sheet[f"{COL_PUT_M5}{row}"] = "❌ Falló PUT"
                    sheet[f"{COL_DET_M5}{row}"] = "Error en subida de archivo"
                    continue

            if script_name:
                cmd_str = ""
                if script_type == "custom":
                    cmd_str = f'runscript -CloudFile="{script_name}"'
                else:
                    cmd_str = f'falconscript -Name="{script_name}"'

                scr_resp = rtr_admin.batch_admin_command(
                    base_command=script_type.replace("custom", "runscript"),
                    batch_id=batch_id,
                    command_string=cmd_str,
                    timeout=CMD_TIMEOUT_SECONDS,
                    timeout_duration=CMD_TIMEOUT_DURATION
                )

                if scr_resp.get("status_code") in (200, 201):
                    sheet[f"{COL_SCRIPT_M5}{row}"] = "✅ Ejecutado"
                    print(f"   {Colors.GREEN}--> {hostname}: Script OK{Colors.RESET}")
                else:
                    sheet[f"{COL_SCRIPT_M5}{row}"] = "❌ Falló Script"

        except KeyboardInterrupt:
            workbook.save(EXCEL_FILE_PATH_MOD5)
            print("\nCancelado por usuario.")
            return
        except Exception as e:
            sheet[f"{COL_DET_M5}{row}"] = f"Error: {str(e)}"

        if done % SAVE_EVERY == 0:
            workbook.save(EXCEL_FILE_PATH_MOD5)

    workbook.save(EXCEL_FILE_PATH_MOD5)
    print(f"\n{Colors.SUCCESS}Despliegue finalizado.{Colors.RESET}")
    input("Presiona Enter para volver...")

# =========================
# MÓDULO 6: UPLOAD SENSORS
# =========================
def get_installers_sensors(base: Path, target_os: str, linux_format: str | None = None) -> list[Path]:
    folder_map = {"windows": base / "Windows", "macos": base / "Mac", "linux": base / "Linux"}
    ext_map = {"windows": {".exe", ".msi"}, "macos": {".pkg"}, "linux": {f".{linux_format}"} if linux_format else {".deb", ".rpm"}}

    folder = folder_map[target_os]
    exts = ext_map[target_os]

    if not folder.exists():
        raise FileNotFoundError(f"No existe la carpeta local: {folder}")

    files = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in exts]
    if not files:
        raise FileNotFoundError(f"Sin instaladores válidos ({exts}) en {folder}")

    return files

def upload_and_put_sensors(rtr: RealTimeResponse, rtr_admin: RealTimeResponseAdmin, aid: str, target_os: str, linux_format: str | None = None) -> None:
    import hashlib  # Necesario para calcular el hash local del nuevo sensor

    installers = get_installers_sensors(BASE_LOCAL_SENSORS, target_os, linux_format)
    remote_dest_dir = DEFAULT_DEST_SENSORS[target_os]

    print(f"\n{Colors.BOLD}{Colors.YELLOW}[i] RECORDATORIO DE ORIGEN:{Colors.RESET}")
    print(f"{Colors.INFO}    Tus instaladores locales se están leyendo desde:{Colors.RESET}")
    print(f"{Colors.INFO}    {BASE_LOCAL_SENSORS}{Colors.RESET}")

    print(f"\n{Colors.INFO}Se encontraron {len(installers)} archivo(s) para {target_os}.{Colors.RESET}")
    print(f"{Colors.INFO}Destino en el servidor: {remote_dest_dir}{Colors.RESET}")

    session_id = None
    try:
        sess = rtr.init_session(device_id=aid)
        require_ok_simple(sess, "init_session")
        session_id = sess["body"]["resources"][0]["session_id"]

        # 1. Asegurar que la ruta exista en el servidor remoto
        ps_mkdir = f"if (!(Test-Path -LiteralPath '{remote_dest_dir}')) {{ New-Item -ItemType Directory -Path '{remote_dest_dir}' -Force | Out-Null }}"
        crid_mkdir = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", f"runscript -Raw=```{ps_mkdir}```")
        wait_admin_command_simple(rtr_admin, crid_mkdir)

        # 2. VALIDAR VERSIONES ANTERIORES EN EL EQUIPO
        print(f"\n{Colors.BLUE}================================================={Colors.RESET}")
        print(f"{Colors.BOLD} VALIDANDO VERSIONES ANTERIORES EN EL EQUIPO{Colors.RESET}")
        print(f"{Colors.BLUE}================================================={Colors.RESET}")
        
        # Script PowerShell que busca estrictamente sensores
        ps_check = (
            "$ErrorActionPreference='SilentlyContinue'; "
            f"$path = '{remote_dest_dir}'; "
            "if (!(Test-Path -LiteralPath $path)) { Write-Output 'NO_DIR'; exit; }; "
            "$files = Get-ChildItem -Path $path -File | Where-Object { $_.Extension -match '\\.(exe|msi|pkg|deb|rpm)$' }; "
            "if (-not $files) { Write-Output 'NO_FILES'; exit; }; "
            "foreach ($f in $files) { "
            "  $hash = (Get-FileHash $f.FullName -Algorithm SHA256).Hash; "
            "  Write-Output \"$($hash)|$($f.Name)\"; "
            "}"
        )
        cmd_check = f"runscript -Raw=```{ps_check}```"
        crid_check = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", cmd_check)
        last_check = wait_admin_command_simple(rtr_admin, crid_check)
        stdout = last_check.get("body", {}).get("resources", [{}])[0].get("stdout", "")

        if "NO_DIR" in stdout:
            print(f"{Colors.WARNING}  [-] La ruta destino aún no existe o está vacía.{Colors.RESET}")
        elif "NO_FILES" in stdout or not stdout.strip():
            print(f"{Colors.INFO}  [i] No hay sensores (.exe, .pkg, .rpm, .deb) almacenados en esta ruta.{Colors.RESET}")
        else:
            lines = stdout.strip().split("\n")
            archivos = [line.split("|", 1) for line in lines if "|" in line]
            if archivos:
                for h, n in archivos:
                    h, n = h.strip(), n.strip()
                    print(f"\n  {Colors.WARNING}SENSOR ENCONTRADO:{Colors.RESET} {n}")
                    print(f"  {Colors.CYAN}Hash SHA256:{Colors.RESET}       {h}")
                    eliminar = input(f"  {Colors.ERROR}¿Deseas ELIMINAR este instalador remoto? (s/n):{Colors.RESET} ").strip().lower()
                    if eliminar in ('s', 'si', 'y', 'yes'):
                        # Sanitizar n (nombre de archivo proveniente del stdout remoto) antes de
                        # interpolarlo en PowerShell: mitiga A1 Injection por nombre malicioso.
                        safe_n = _sanitize_ps_literal(n)
                        full_path = f"{remote_dest_dir}\\{safe_n}"
                        del_ps = f"Remove-Item -LiteralPath '{full_path}' -Force"
                        del_cmd = f"runscript -Raw=```{del_ps}```"
                        crid_del = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", del_cmd)
                        wait_admin_command_simple(rtr_admin, crid_del)
                        print(f"  {Colors.SUCCESS}[OK] Archivo eliminado del servidor.{Colors.RESET}")
                    else:
                        print(f"  {Colors.INFO}[-] Archivo conservado en el servidor.{Colors.RESET}")

        # 3. SUBIR EL NUEVO SENSOR
        print(f"\n{Colors.BLUE}================================================={Colors.RESET}")
        print(f"{Colors.BOLD} PREPARANDO Y CARGANDO NUEVO SENSOR{Colors.RESET}")
        print(f"{Colors.BLUE}================================================={Colors.RESET}")

        crid_cd = rtr_admin_cmd_simple(rtr_admin, session_id, "cd", f'cd "{remote_dest_dir}"')
        wait_admin_command_simple(rtr_admin, crid_cd)

        for installer in installers:
            # Calcular el Hash Local del nuevo sensor
            h = hashlib.sha256()
            with open(installer, 'rb') as f_hash:
                for chunk in iter(lambda: f_hash.read(4096), b""):
                    h.update(chunk)
            local_hash = h.hexdigest().upper()

            # Evitar el Error 40006 de CrowdStrike eliminando espacios del nombre
            safe_name = installer.name.replace(" ", "_").replace("-", "_")
            put_name = f"{Path(safe_name).stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{installer.suffix}"
            
            print(f"\n{Colors.CYAN}---> Procesando nuevo archivo:{Colors.RESET} {installer.name}")
            print(f"{Colors.YELLOW}     Hash SHA256 (NUEVO):{Colors.RESET}      {local_hash}")

            with open(installer, "rb") as f:
                resp_put = rtr_admin.create_put_files(
                    name=put_name,
                    description=f"Auto-upload: {installer.name}",
                    comments_for_audit_log="RTR upload",
                    files=[('file', (put_name, f.read(), "application/octet-stream"))]
                )
            require_ok_simple(resp_put, "create_put_files")
            print(f"  {Colors.SUCCESS}[OK] Subido a repositorio Put Files (CrowdStrike Cloud){Colors.RESET}")

            # Se añaden comillas por seguridad al enviarlo al equipo
            crid_put = rtr_admin_cmd_simple(rtr_admin, session_id, "put", f'put "{put_name}"')
            wait_admin_command_simple(rtr_admin, crid_put)
            print(f"  {Colors.SUCCESS}[OK] Depositado exitosamente en el servidor (Ruta Destino){Colors.RESET}")

            time.sleep(1)

    finally:
        if session_id:
            try:
                rtr.delete_session(session_id=session_id)
            except Exception:
                pass

def manage_remote_sensors(rtr: RealTimeResponse, rtr_admin: RealTimeResponseAdmin, aid: str) -> None:
    print(f"\n{Colors.BOLD}{Colors.YELLOW}[i] RECORDATORIO IMPORTANTE:{Colors.RESET}")
    print(f"{Colors.INFO}    El origen de los instaladores en tu equipo es:{Colors.RESET}")
    print(f"{Colors.INFO}    {BASE_LOCAL_SENSORS}{Colors.RESET}")
    
    session_id = None
    try:
        sess = rtr.init_session(device_id=aid)
        require_ok_simple(sess, "init_session")
        session_id = sess["body"]["resources"][0]["session_id"]
        
        for os_name, remote_path in DEFAULT_DEST_SENSORS.items():
            print(f"\n{Colors.BLUE}{'='*60}{Colors.RESET}")
            print(f"{Colors.BOLD} Explorando ruta {os_name.upper()}:{Colors.RESET} {remote_path}")
            print(f"{Colors.BLUE}{'='*60}{Colors.RESET}")
            
            ps = (
                "$ErrorActionPreference='SilentlyContinue'; "
                f"if (!(Test-Path -LiteralPath '{remote_path}')) {{ Write-Output 'NO_DIR'; exit; }}; "
                f"$files = Get-ChildItem -Path '{remote_path}' -File; "
                "if ($files.Count -eq 0) { Write-Output 'NO_FILES'; exit; }; "
                "foreach ($f in $files) { "
                "  $hash = (Get-FileHash $f.FullName -Algorithm SHA256).Hash; "
                "  Write-Output \"$($hash)|$($f.Name)\"; "
                "}"
            )
            cmd_runscript = f"runscript -Raw=```{ps}```"
            
            crid = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", cmd_runscript)
            last = wait_admin_command_simple(rtr_admin, crid)
            
            stdout = last.get("body", {}).get("resources", [{}])[0].get("stdout", "")
            
            if "NO_DIR" in stdout:
                print(f"{Colors.WARNING}  [-] La carpeta no existe en este servidor.{Colors.RESET}")
                continue
            elif "NO_FILES" in stdout or not stdout.strip():
                print(f"{Colors.INFO}  [i] No hay sensores almacenados en esta ruta.{Colors.RESET}")
                continue
                
            lines = stdout.strip().split("\n")
            archivos = []
            for line in lines:
                if "|" in line:
                    h, n = line.split("|", 1)
                    archivos.append((h.strip(), n.strip()))
                    
            if not archivos:
                print(f"{Colors.WARNING}  [-] Error leyendo archivos de la ruta.{Colors.RESET}")
                continue
                
            for h, n in archivos:
                print(f"\n  {Colors.CYAN}Archivo: {Colors.RESET}{n}")
                print(f"  {Colors.CYAN}Hash:    {Colors.RESET}{h}")
                
                eliminar = input(f"  {Colors.WARNING}¿Deseas eliminar este instalador remoto? (s/n):{Colors.RESET} ").strip().lower()
                if eliminar in ('s', 'si', 'y', 'yes'):
                    # Sanitizar n (nombre de archivo proveniente del stdout remoto) antes de
                    # interpolarlo en PowerShell: mitiga A1 Injection por nombre malicioso.
                    safe_n = _sanitize_ps_literal(n)
                    full_path = f"{remote_path}\\{safe_n}"
                    del_ps = f"Remove-Item -LiteralPath '{full_path}' -Force"
                    del_cmd = f"runscript -Raw=```{del_ps}```"
                    crid_del = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", del_cmd)
                    wait_admin_command_simple(rtr_admin, crid_del)
                    print(f"  {Colors.SUCCESS}[OK] Archivo eliminado del servidor.{Colors.RESET}")
                else:
                    print(f"  {Colors.INFO}[-] Archivo conservado en el servidor.{Colors.RESET}")
                    
    finally:
        if session_id:
            try:
                rtr.delete_session(session_id=session_id)
            except:
                pass

def run_upload_sensors_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("Carga de Sensores (Windows/Mac/Linux)")

    entrada = input(f"\n{Colors.CYAN}AID o Hostname del destino (ej. MDEOLYFS01):{Colors.RESET} ").strip()
    if not entrada:
        print(f"{Colors.ERROR}Entrada vacía. Abortando.{Colors.RESET}")
        return

    rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
    validar_credenciales_falconpy(rtr)
    rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)

    from falconpy import Hosts
    hosts = Hosts(auth_object=rtr.auth_object)
    
    aid = entrada
    if len(entrada) != 32:
        print(f"{Colors.INFO}Buscando el AID para el hostname '{entrada}'...{Colors.RESET}")
        resp = hosts.query_devices_by_filter_scroll(filter=f"hostname:'{entrada}'")
        resources = resp.get("body", {}).get("resources") or []
        if not resources:
            print(f"{Colors.ERROR}No se encontró ningún equipo con el hostname '{entrada}' en la consola.{Colors.RESET}")
            input(f"\n{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
            return
        aid = resources[0]
        print(f"{Colors.SUCCESS}¡Equipo encontrado! AID: {aid}{Colors.RESET}")

    while True:
        print(f"\n{Colors.BLUE}==============================================={Colors.RESET}")
        print(f"{Colors.BOLD} MENÚ DE GESTIÓN Y CARGA DE SENSORES{Colors.RESET}")
        print(f"{Colors.BLUE}==============================================={Colors.RESET}")
        print(f"{Colors.INFO}1){Colors.RESET} Cargar Sensor para Windows (.exe / .msi)")
        print(f"{Colors.INFO}2){Colors.RESET} Cargar Sensor para Mac (.pkg)")
        print(f"{Colors.INFO}3){Colors.RESET} Cargar Sensores para Linux (.deb o .rpm)")
        print(f"{Colors.INFO}4){Colors.RESET} Gestionar versiones remotas (Listar, Hashes, Eliminar)")
        print(f"{Colors.INFO}5){Colors.RESET} Volver al menú principal")
        print(f"{Colors.BLUE}==============================================={Colors.RESET}")

        opcion = input(f"{Colors.CYAN}Elige una opción (1-5):{Colors.RESET} ").strip()

        try:
            if opcion == "1":
                upload_and_put_sensors(rtr, rtr_admin, aid, target_os="windows")
            elif opcion == "2":
                upload_and_put_sensors(rtr, rtr_admin, aid, target_os="macos")
            elif opcion == "3":
                linux_format = input(f"{Colors.CYAN}¿Formato Linux? (deb/rpm) [default: deb]:{Colors.RESET} ").strip().lower()
                if not linux_format:
                    linux_format = "deb"
                upload_and_put_sensors(rtr, rtr_admin, aid, target_os="linux", linux_format=linux_format)
            elif opcion == "4":
                manage_remote_sensors(rtr, rtr_admin, aid)
            elif opcion == "5":
                print(f"{Colors.INFO}Volviendo...{Colors.RESET}")
                break
            else:
                print(f"{Colors.ERROR}Opción no válida. Intenta de nuevo.{Colors.RESET}")
        except Exception as e:
            print(f"{Colors.ERROR}Error durante la operación: {e}{Colors.RESET}")

# =========================
# MÓDULO 7: VALIDAR HASH EN VIRUSTOTAL
# =========================
def consultar_hash_virustotal(hash_str: str) -> str:
    url = f"https://www.virustotal.com/api/v3/files/{hash_str}"
    headers = {"x-apikey": VIRUSTOTAL_API_KEY}
    try:
        r = requests.get(url, headers=headers)
        r.raise_for_status()
        stats = r.json()["data"]["attributes"]["last_analysis_stats"]
        maliciosos = stats.get("malicious", 0)
        sospechosos = stats.get("suspicious", 0)

        if maliciosos > 0 or sospechosos > 0:
            return f"Detectado: {maliciosos} maliciosos, {sospechosos} sospechosos"
        else:
            return "Limpio: 0 detecciones"

    except requests.exceptions.HTTPError as e:
        if r.status_code == 404:
            return "No se encuentran registros maliciosos o no encontrado en VirusTotal"
        elif r.status_code == 401:
            return "Error: Clave API inválida o sin permisos"
        elif r.status_code == 429:
            return "Error: Límite de cuota de API excedido"
        else:
            return f"Error HTTP {r.status_code}"
    except Exception as e:
        return f"Error de conexión: {e}"

def run_virustotal_hash_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("RTR -> VirusTotal (C:\\proc_ent)")

    aid = input(f"{Colors.CYAN}AID del host (Host ID):{Colors.RESET} " ).strip()
    if not aid:
        print(f"{Colors.ERROR}AID vacío. Abortando.{Colors.RESET}")
        return

    print(f"\n{Colors.INFO}Conectando a RTR...{Colors.RESET}")
    rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
    validar_credenciales_falconpy(rtr)
    rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)

    session_id = None
    try:
        sess = rtr.init_session(device_id=aid)
        require_ok_simple(sess, "init_session")
        session_id = sess["body"]["resources"][0]["session_id"]
        print(f"{Colors.SUCCESS}[OK]{Colors.RESET} RTR session_id: {session_id}")

        remote_dir = r"C:\\proc_ent"
        ps = (
            "$ErrorActionPreference='SilentlyContinue'; "
            f"if (!(Test-Path -LiteralPath '{remote_dir}')) {{ Write-Output 'NO_DIR'; exit; }}; "
            f"$files = Get-ChildItem -Path '{remote_dir}' -File; "
            "if ($files.Count -eq 0) { Write-Output 'NO_FILES'; exit; }; "
            "foreach ($f in $files) { "
            "  $hash = (Get-FileHash $f.FullName -Algorithm SHA256).Hash; "
            "  Write-Output \"$($hash)|$($f.Name)\"; "
            "}"
        )
        cmd_runscript = f"runscript -Raw=```{ps}```"

        print(f"{Colors.INFO}Calculando hashes remotamente en {remote_dir}...{Colors.RESET}")
        crid = rtr_admin_cmd_simple(rtr_admin, session_id, "runscript", cmd_runscript)
        last = wait_admin_command_simple(rtr_admin, crid)

        stdout = last.get("body", {}).get("resources", [{}])[0].get("stdout", "")

        if "NO_DIR" in stdout:
            print(f"{Colors.WARNING}[WARN]{Colors.RESET} El directorio {remote_dir} no existe en este host.")
            return
        elif "NO_FILES" in stdout:
            print(f"{Colors.WARNING}[WARN]{Colors.RESET} No hay archivos en {remote_dir}.")
            return

        lines = stdout.strip().split("\n")
        archivos_hashes = []
        for line in lines:
            if "|" in line:
                h, n = line.split("|", 1)
                archivos_hashes.append((h.strip(), n.strip()))

        if not archivos_hashes:
            print(f"{Colors.WARNING}[WARN]{Colors.RESET} No se pudieron extraer hashes (stdout: {stdout})")
            return

        print(f"\n{Colors.INFO}Consultando {len(archivos_hashes)} archivo(s) en VirusTotal...{Colors.RESET}\n")

        for h, n in archivos_hashes:
            print(f"{Colors.CYAN}Archivo:{Colors.RESET} {n}")
            resultado = consultar_hash_virustotal(h)
            if "Detectado" in resultado or "Error" in resultado:
                color = Colors.RED
            elif "No encontrado" in resultado:
                color = Colors.YELLOW
            else:
                color = Colors.GREEN
            print(f"  -> Hash: {h}")
            print(f"  -> VT: {color}{resultado}{Colors.RESET}\n")

    finally:
        if session_id:
            try:
                rtr.delete_session(session_id=session_id)
            except Exception:
                pass


# =========================
# MÓDULO 8: DELETE HASH
# =========================
def run_delete_hash_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("Delete Hash RTR")

    def first_resource(resp):
        body = resp.get("body", {}) if isinstance(resp, dict) else {}
        resources = body.get("resources", [])
        if isinstance(resources, list) and resources:
            return resources[0]
        return {}

    def extract_errors(resp):
        if not isinstance(resp, dict):
            return str(resp)
        body = resp.get("body", {})
        errors = body.get("errors", [])
        if errors:
            msgs = []
            for err in errors:
                if isinstance(err, dict):
                    msgs.append(err.get("message") or str(err))
                else:
                    msgs.append(str(err))
            return " | ".join(msgs)
        return str(body) if body else str(resp)

    def split_csv(text):
        return [x.strip() for x in text.split(",") if x.strip()]

    def normalize_hashes(raw_hashes):
        valid = set()
        invalid = []
        for h in raw_hashes:
            h = h.strip().lower()
            if re.fullmatch(r"[a-f0-9]{64}", h):
                valid.add(h)
            else:
                invalid.append(h)
        return valid, invalid

    def dedupe_keep_order(items):
        seen = set()
        out = []
        for item in items:
            key = item.strip().lower()
            if key not in seen:
                seen.add(key)
                out.append(item)
        return out

    def normalize_windows_path(path):
        p = path.strip().strip('"').strip("'")
        p = p.replace("/", "\\")
        p = re.sub(r"\\{2,}", r"\\", p)
        return p

    def wait_for_readonly_result(rtr, cloud_request_id, max_wait=120, sleep_s=2):
        deadline = time.time() + max_wait
        last_resp = None
        while time.time() < deadline:
            resp = rtr.check_command_status(
                cloud_request_id=cloud_request_id,
                sequence_id=0
            )
            last_resp = resp
            if resp.get("status_code") == 200:
                res = first_resource(resp)
                if res.get("complete") is True:
                    return resp
                if res.get("stdout") or res.get("stderr"):
                    return resp
            time.sleep(sleep_s)
        return last_resp

    def wait_for_admin_result(rtr_admin, cloud_request_id, max_wait=120, sleep_s=2):
        deadline = time.time() + max_wait
        last_resp = None
        while time.time() < deadline:
            resp = rtr_admin.check_admin_command_status(
                cloud_request_id=cloud_request_id,
                sequence_id=0
            )
            last_resp = resp
            if resp.get("status_code") == 200:
                res = first_resource(resp)
                if res.get("complete") is True:
                    return resp
                if res.get("stdout") or res.get("stderr"):
                    return resp
            time.sleep(sleep_s)
        return last_resp

    def run_readonly_command(rtr, session_id, base_command, command_string):
        submit = rtr.execute_command(
            session_id=session_id,
            base_command=base_command,
            command_string=command_string
        )
        if submit.get("status_code") not in (200, 201):
            return False, "", extract_errors(submit), submit

        req = first_resource(submit)
        cloud_request_id = req.get("cloud_request_id")
        if not cloud_request_id:
            return True, req.get("stdout", ""), req.get("stderr", ""), submit

        result = wait_for_readonly_result(rtr, cloud_request_id)
        if not result or result.get("status_code") != 200:
            return False, "", extract_errors(result or submit), result or submit

        res = first_resource(result)
        return True, res.get("stdout", "") or "", res.get("stderr", "") or "", result

    def run_admin_command(rtr_admin, session_id, base_command, command_string):
        submit = rtr_admin.execute_admin_command(
            session_id=session_id,
            base_command=base_command,
            command_string=command_string
        )
        if submit.get("status_code") not in (200, 201):
            return False, "", extract_errors(submit), submit

        req = first_resource(submit)
        cloud_request_id = req.get("cloud_request_id")
        if not cloud_request_id:
            return True, req.get("stdout", ""), req.get("stderr", ""), submit

        result = wait_for_admin_result(rtr_admin, cloud_request_id)
        if not result or result.get("status_code") != 200:
            return False, "", extract_errors(result or submit), result or submit

        res = first_resource(result)
        return True, res.get("stdout", "") or "", res.get("stderr", "") or "", result

    def build_powershell_discovery(target_path):
        target_path = normalize_windows_path(target_path)
        escaped = target_path.replace("'", "''")
        ps = rf'''$ErrorActionPreference = "Stop"
$target = '{escaped}'
$result = @()

if (Test-Path -LiteralPath $target) {{
    $item = Get-Item -LiteralPath $target -Force
    if ($item.PSIsContainer) {{
        $result = Get-ChildItem -LiteralPath $target -Recurse -File -Force -ErrorAction SilentlyContinue |
            Select-Object -ExpandProperty FullName
    }} else {{
        $result = @($item.FullName)
    }}
}}

$result | ConvertTo-Json -Compress
'''.strip()
        return ps

    def run_ps_discovery(rtr_admin, session_id, target_path):
        target_path = normalize_windows_path(target_path)
        ps = build_powershell_discovery(target_path)
        cmd = f"runscript -Raw=```{ps}```"

        ok, stdout, stderr, raw = run_admin_command(
            rtr_admin,
            session_id,
            base_command="runscript",
            command_string=cmd
        )
        if not ok:
            return False, [], stderr or extract_errors(raw)

        output = "\n".join([x for x in [stdout, stderr] if x]).strip()
        if not output:
            return True, [], ""

        json_match = re.search(r"(\[.*\]|\"[^\"]+\")", output, re.S)
        if not json_match:
            return False, [], output

        try:
            parsed = json.loads(json_match.group(1))
            if isinstance(parsed, str):
                return True, [normalize_windows_path(parsed)], ""
            if isinstance(parsed, list):
                return True, [normalize_windows_path(x) for x in parsed if x], ""
            return True, [], ""
        except Exception:
            return False, [], output

    def extract_sha256(text):
        if not text:
            return None
        m = re.search(r"(?i)sha256[^a-f0-9]*([a-f0-9]{64})", text)
        if m:
            return m.group(1).lower()
        m = re.search(r"\b([a-f0-9]{64})\b", text, re.I)
        if m:
            return m.group(1).lower()
        return None

    def text_has_error(text):
        if not text:
            return False
        t = text.lower()
        indicators = [
            "error", "failed", "no such file", "cannot find",
            "permission denied", "access is denied",
            "you can't hash a folder", "does not exist",
            "command not found", "path not found", "file not found",
            "remove-item", "unauthorized", "not permitted", "in use",
            "being used by another process", "sharing violation",
            "couldn't find"
        ]
        return any(x in t for x in indicators)

    def verify_file_deleted(rtr, session_id, file_path):
        file_path = normalize_windows_path(file_path)

        ok_hash, stdout_hash, stderr_hash, raw_hash = run_readonly_command(
            rtr,
            session_id,
            base_command="filehash",
            command_string=f'filehash "{file_path}"'
        )

        hash_output = "\n".join([x for x in [stdout_hash, stderr_hash] if x]).strip()

        if (not ok_hash) or text_has_error(hash_output):
            return True, hash_output or extract_errors(raw_hash)

        sha256 = extract_sha256(hash_output)
        if sha256:
            return False, f"El archivo sigue existiendo. SHA256 actual: {sha256}"

        return False, hash_output or "No se pudo confirmar el borrado."

    def run_powershell_delete(rtr_admin, session_id, file_path):
        file_path = normalize_windows_path(file_path)
        escaped = file_path.replace("'", "''")

        ps = rf'''$ErrorActionPreference = "Stop"
$p = '{escaped}'

Write-Output ("RAW_PATH=" + $p)

if (-not (Test-Path -LiteralPath $p)) {{
    Write-Output "NOT_FOUND"
    exit 0
}}

$item = Get-Item -LiteralPath $p -Force
Write-Output ("TARGET=" + $item.FullName)
Write-Output ("ATTRIBUTES=" + $item.Attributes)

try {{
    Remove-Item -LiteralPath $p -Force -ErrorAction Stop
}} catch {{
    Write-Output ("DELETE_ERROR=" + $_.Exception.Message)
    exit 1
}}

if (Test-Path -LiteralPath $p) {{
    Write-Output "STILL_EXISTS"
    exit 2
}} else {{
    Write-Output "DELETED"
    exit 0
}}
'''.strip()

        cmd = f"runscript -Raw=```{ps}```"

        ok, stdout, stderr, raw = run_admin_command(
            rtr_admin,
            session_id,
            base_command="runscript",
            command_string=cmd
        )

        output = "\n".join([x for x in [stdout, stderr] if x]).strip()

        if not ok:
            return False, output or extract_errors(raw)

        if "DELETED" in output or "NOT_FOUND" in output:
            return True, output

        return False, output or "No se pudo confirmar el borrado por PowerShell."

    def run_delete_and_verify(rtr, rtr_admin, session_id, file_path):
        file_path = normalize_windows_path(file_path)

        ps_ok, ps_detail = run_powershell_delete(rtr_admin, session_id, file_path)
        if ps_ok:
            deleted_after_ps, verify_after_ps = verify_file_deleted(rtr, session_id, file_path)
            if deleted_after_ps:
                return True, "Borrado confirmado con PowerShell LiteralPath."
            return False, verify_after_ps or ps_detail

        return False, ps_detail

    def ask_delete_confirmation(aid, matched_files):
        print(f"\n{Colors.WARNING}[?] Archivos validados para borrar en host {aid}:{Colors.RESET}")
        for mf in matched_files:
            print(f"    - {mf}")

        answer = input(
            f"\n    {Colors.ERROR}¿Deseas eliminarlos ahora en este mismo host? (s/n): {Colors.RESET}"
        ).strip().lower()

        return answer in ("s", "si", "sí", "y", "yes")

    print(f"{Colors.BLUE}[*] Autenticando con CrowdStrike...{Colors.RESET}")
    try:
        rtr = RealTimeResponse(client_id=client_id, client_secret=client_secret)
        validar_credenciales_falconpy(rtr)
        rtr_admin = RealTimeResponseAdmin(auth_object=rtr.auth_object)
    except Exception as e:
        print(f"{Colors.ERROR}Error creando clientes RTR: {e}{Colors.RESET}")
        input(f"{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
        return

    print(f"\n{Colors.BLUE}[*] Información del Objetivo{Colors.RESET}")
    aid_input = input(f"    {Colors.CYAN}Ingresa los Device IDs (AIDs) separados por coma: {Colors.RESET}").strip()
    aids = split_csv(aid_input)
    if not aids:
        print(f"{Colors.ERROR}Debes ingresar al menos un AID.{Colors.RESET}")
        input(f"{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
        return

    print(f"\n{Colors.BLUE}[*] Hashes SHA256 permitidos para depuración{Colors.RESET}")
    hash_input = input(f"    {Colors.CYAN}Ingresa los SHA256 separados por coma: {Colors.RESET}").strip()
    valid_hashes, invalid_hashes = normalize_hashes(split_csv(hash_input))
    if invalid_hashes:
        print(f"{Colors.ERROR}Hashes inválidos detectados:{Colors.RESET}")
        for h in invalid_hashes:
            print(f"    - {h}")
        input(f"{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
        return
    if not valid_hashes:
        print(f"{Colors.ERROR}Debes ingresar al menos un SHA256 válido.{Colors.RESET}")
        input(f"{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
        return

    print(f"\n{Colors.BLUE}[*] Rutas a revisar{Colors.RESET}")
    print(f"{Colors.INFO}    Puedes ingresar archivos o carpetas.{Colors.RESET}")
    print(f"{Colors.INFO}    Primero valida hashes y luego te preguntará si deseas borrar.{Colors.RESET}")
    print(f"    {Colors.WARNING}(Escribe 'FIN' y presiona Enter cuando termines){Colors.RESET}")

    targets = []
    c = 1
    while True:
        ruta = input(f"    Ruta {c}: ").strip()
        if ruta.upper() == "FIN" or ruta == "":
            break
        targets.append(normalize_windows_path(ruta))
        c += 1

    if not targets:
        print(f"{Colors.ERROR}No ingresaste rutas.{Colors.RESET}")
        input(f"{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")
        return

    resumen = {
        "hosts_ok": 0,
        "hosts_error": 0,
        "targets_sin_archivos": 0,
        "files_evaluated": 0,
        "files_matching_hash": 0,
        "files_deleted": 0,
        "files_skipped_hash": 0,
        "files_hash_error": 0,
        "files_delete_failed": 0,
        "delete_cancelled": 0
    }

    for aid in aids:
        print(f"\n{Colors.BLUE}================================================={Colors.RESET}")
        print(f"{Colors.INFO}[+] Conectando al equipo: {aid}{Colors.RESET}")

        init = rtr.init_session(device_id=aid)
        if init.get("status_code") != 201:
            print(f"{Colors.ERROR}  [-] No se pudo iniciar sesión RTR: {extract_errors(init)}{Colors.RESET}")
            resumen["hosts_error"] += 1
            continue

        session_id = first_resource(init).get("session_id")
        if not session_id:
            print(f"{Colors.ERROR}  [-] No vino session_id en la respuesta.{Colors.RESET}")
            resumen["hosts_error"] += 1
            continue

        resumen["hosts_ok"] += 1
        print(f"{Colors.SUCCESS}  [+] Sesión RTR iniciada correctamente.{Colors.RESET}")

        matched_files = []

        try:
            for target in targets:
                print(f"\n{Colors.WARNING}  [*] Explorando: {target}{Colors.RESET}")

                ok_disc, files_found, disc_error = run_ps_discovery(rtr_admin, session_id, target)
                if not ok_disc:
                    print(f"{Colors.ERROR}    [X] No se pudo listar el contenido/ruta.{Colors.RESET}")
                    if disc_error:
                        print(f"{Colors.ERROR}        Detalle: {disc_error}{Colors.RESET}")
                    resumen["targets_sin_archivos"] += 1
                    continue

                if not files_found:
                    print(f"{Colors.WARNING}    [-] No se encontraron archivos en esa ruta.{Colors.RESET}")
                    resumen["targets_sin_archivos"] += 1
                    continue

                print(f"{Colors.INFO}    [i] Archivos encontrados: {len(files_found)}{Colors.RESET}")

                for file_path in files_found:
                    file_path = normalize_windows_path(file_path)
                    resumen["files_evaluated"] += 1
                    print(f"{Colors.WARNING}    [*] Hashing: {file_path}{Colors.RESET}")

                    ok_hash, stdout_hash, stderr_hash, raw_hash = run_readonly_command(
                        rtr,
                        session_id,
                        base_command="filehash",
                        command_string=f'filehash "{file_path}"'
                    )

                    hash_output = "\n".join([x for x in [stdout_hash, stderr_hash] if x]).strip()
                    if (not ok_hash) or text_has_error(hash_output):
                        print(f"{Colors.ERROR}      [X] No se pudo calcular hash.{Colors.RESET}")
                        detalle = hash_output or extract_errors(raw_hash)
                        if detalle:
                            print(f"{Colors.ERROR}          Detalle: {detalle}{Colors.RESET}")
                        resumen["files_hash_error"] += 1
                        continue

                    sha256 = extract_sha256(hash_output)
                    if not sha256:
                        print(f"{Colors.ERROR}      [X] No se pudo extraer SHA256.{Colors.RESET}")
                        if hash_output:
                            print(f"{Colors.ERROR}          Salida: {hash_output}{Colors.RESET}")
                        resumen["files_hash_error"] += 1
                        continue

                    print(f"{Colors.INFO}      [i] SHA256: {sha256}{Colors.RESET}")

                    if sha256 not in valid_hashes:
                        print(f"{Colors.WARNING}      [-] Hash no autorizado. No se borra.{Colors.RESET}")
                        resumen["files_skipped_hash"] += 1
                        continue

                    print(f"{Colors.SUCCESS}      [MATCH] Coincide con lista autorizada.{Colors.RESET}")
                    matched_files.append(file_path)
                    resumen["files_matching_hash"] += 1

            matched_files = dedupe_keep_order(matched_files)

            if matched_files:
                should_delete = ask_delete_confirmation(aid, matched_files)

                if should_delete:
                    print(f"\n{Colors.WARNING}[!] Ejecutando borrado en host {aid}...{Colors.RESET}")
                    for file_path in matched_files:
                        ok_delete, detail = run_delete_and_verify(
                            rtr,
                            rtr_admin,
                            session_id,
                            file_path
                        )

                        if ok_delete:
                            print(f"{Colors.SUCCESS}    [✓] BORRADO CONFIRMADO: {file_path}{Colors.RESET}")
                            resumen["files_deleted"] += 1
                        else:
                            print(f"{Colors.ERROR}    [X] NO BORRADO: {file_path}{Colors.RESET}")
                            if detail:
                                print(f"{Colors.ERROR}        Detalle: {detail}{Colors.RESET}")
                            resumen["files_delete_failed"] += 1

                        time.sleep(1)
                else:
                    print(f"{Colors.WARNING}    [-] Borrado cancelado por el usuario para el host {aid}.{Colors.RESET}")
                    resumen["delete_cancelled"] += len(matched_files)
            else:
                print(f"\n{Colors.WARNING}  [-] No hubo coincidencias para borrar en este host.{Colors.RESET}")

        finally:
            print(f"\n{Colors.INFO}  [+] Cerrando sesión RTR...{Colors.RESET}")
            try:
                rtr.delete_session(session_id=session_id)
            except Exception:
                pass

    print(f"\n{Colors.BOLD}{Colors.CYAN}================ RESUMEN ================{Colors.RESET}")
    print(f"{Colors.SUCCESS}Hosts con sesión RTR: {resumen['hosts_ok']}{Colors.RESET}")
    print(f"{Colors.ERROR}Hosts con error de sesión: {resumen['hosts_error']}{Colors.RESET}")
    print(f"{Colors.WARNING}Targets sin archivos encontrados: {resumen['targets_sin_archivos']}{Colors.RESET}")
    print(f"{Colors.RESULT}Archivos evaluados: {resumen['files_evaluated']}{Colors.RESET}")
    print(f"{Colors.SUCCESS}Coincidencias por hash: {resumen['files_matching_hash']}{Colors.RESET}")
    print(f"{Colors.SUCCESS}Archivos borrados: {resumen['files_deleted']}{Colors.RESET}")
    print(f"{Colors.WARNING}Archivos omitidos por hash distinto: {resumen['files_skipped_hash']}{Colors.RESET}")
    print(f"{Colors.WARNING}Archivos con error de hash: {resumen['files_hash_error']}{Colors.RESET}")
    print(f"{Colors.ERROR}Intentos de borrado fallidos: {resumen['files_delete_failed']}{Colors.RESET}")
    print(f"{Colors.WARNING}Coincidencias no borradas por cancelación: {resumen['delete_cancelled']}{Colors.RESET}")

    input(f"\n{Colors.CYAN}Listo. Presione Enter para volver al menú...{Colors.RESET}")

# =========================
# MÓDULO 9: DEPURADOR DE HOSTS
# =========================
def hide_hosts_batch(id_list, sdk: Hosts):
    """Oculta un lote de hosts. Falcon soporta hasta 100 por petición."""
    res = sdk.perform_action(action_name="hide_host", ids=id_list)
    if res["status_code"] != 202:
        for err in res.get("body", {}).get("errors", []):
            print(f"{Colors.ERROR}   Error [{err.get('code')}]: {err.get('message')}{Colors.RESET}")

def hide_hosts(host_list: list, sdk: Hosts):
    """Divide la lista en lotes de 100 y usa ThreadPoolExecutor para ocultarlos en paralelo."""
    batches = [host_list[i:i+100] for i in range(0, len(host_list), 100)]
    with ThreadPoolExecutor() as executor:
        # Forzar la ejecución del map para que no sea un generador perezoso
        list(executor.map(lambda batch: hide_hosts_batch(batch, sdk), batches))

def run_depurador_module(client_id: str, client_secret: str) -> None:
    clear_screen()
    print_banner("Depuración de Hosts Duplicados en Consola")
    
    # Directorio de salida
    ruta_carpeta = Path(f"Depuracion_{datetime.now().strftime('%Y-%m-%d')}")
    ruta_carpeta.mkdir(parents=True, exist_ok=True)
    
    print(f"\n{Colors.INFO}Iniciando módulo de depuración con credenciales en memoria...{Colors.RESET}")
    falcon = Hosts(client_id=client_id, client_secret=client_secret)
    
    print(f"\n{Colors.BLUE}[+] Consultando inventario inicial en la consola...{Colors.RESET}")
    try:
        initial_query = falcon.query_devices_by_filter_scroll(limit=1)
        total = initial_query["body"]["meta"]["pagination"]["total"]
    except Exception as e:
        print(f"{Colors.ERROR}[-] Error al conectar con la API de Falcon. Verifique credenciales. Detalle: {e}{Colors.RESET}")
        return

    offset = ''
    data = []

    with tqdm(total=total, desc="Descargando hosts", unit="host", 
              bar_format="{l_bar}%s{bar}%s{r_bar}" % (Colors.CYAN, Colors.RESET)) as pbar:
        while total > 0:
            response = falcon.query_devices_by_filter_scroll(offset=offset, limit=5000)
            
            if "body" not in response or "meta" not in response["body"]:
                break
                
            current_count = len(response["body"].get("resources", []))
            if current_count == 0:
                break
                
            pbar.update(current_count)
            total -= current_count
            offset = response["body"]["meta"]["pagination"]["offset"]
            
            device_ids = response["body"]["resources"]
            detail_response = falcon.get_device_details(ids=device_ids)
            
            for detail in detail_response["body"].get("resources", []):
                sensor_tags = str(detail.get("tags", "")).replace(",", ";")
                
                # Exclusión de servidores SOX por seguridad
                if 'SOX' in sensor_tags.upper():
                    continue
                        
                data.append({
                    "nombre": detail.get("hostname", "Not Found"),
                    "last_seen": detail.get("last_seen"),
                    "platform": detail.get("platform_name", "Not Found"),
                    "host_id": detail.get("device_id", "Not Found"),
                    "serial": detail.get("serial_number", "Not Found"),
                    "sensor_tags": sensor_tags
                })

    df = pd.DataFrame(data)
    if df.empty:
        print(f"\n{Colors.WARNING}[-] No se encontraron hosts para procesar.{Colors.RESET}")
        return

    excel_inventario = ruta_carpeta / "Inventario_Inicial.xlsx"
    df.to_excel(excel_inventario, index=False)
    print(f"\n{Colors.SUCCESS}[+] Inventario inicial ({len(df)} hosts) guardado en {ruta_carpeta.name}/Inventario_Inicial.xlsx{Colors.RESET}")

    dep_name = pd.DataFrame()
    dep_serial = pd.DataFrame()
    hosts_a_ocultar = set()

    # Procesamiento de fechas para Pandas
    df['last_seen'] = pd.to_datetime(df['last_seen'], utc=True, errors='coerce').dt.tz_localize(None)

    # Duplicados por Nombre
    print(f"\n{Colors.BLUE}[+] Analizando duplicados por Hostname...{Colors.RESET}")
    df_valid_names = df[~df['nombre'].isin(['0', 'Not Found', ''])]
    if not df_valid_names.empty:
        df_sorted = df_valid_names.sort_values(by=['nombre', 'last_seen'], ascending=[True, False])
        dep_name = df_sorted[df_sorted.duplicated(subset=['nombre'], keep='first')]
        if not dep_name.empty:
            print(f"{Colors.WARNING}    -> Encontrados {len(dep_name)} hosts antiguos duplicados por nombre.{Colors.RESET}")
            hosts_a_ocultar.update(dep_name['host_id'].tolist())
        else:
            print(f"{Colors.INFO}    -> No hay duplicados por nombre.{Colors.RESET}")

    # Duplicados por Serial
    print(f"{Colors.BLUE}[+] Analizando duplicados por Serial...{Colors.RESET}")
    df_valid_serials = df[~df['serial'].isin(['0', 'Not Found', '', 'TO BE FILLED BY O.E.M.', 'DEFAULT STRING'])]
    if not df_valid_serials.empty:
        df_sorted = df_valid_serials.sort_values(by=['serial', 'last_seen'], ascending=[True, False])
        dep_serial = df_sorted[df_sorted.duplicated(subset=['serial'], keep='first')]
        if not dep_serial.empty:
            print(f"{Colors.WARNING}    -> Encontrados {len(dep_serial)} hosts antiguos duplicados por serial.{Colors.RESET}")
            hosts_a_ocultar.update(dep_serial['host_id'].tolist())
        else:
            print(f"{Colors.INFO}    -> No hay duplicados por serial.{Colors.RESET}")

    # Ocultar Hosts (Papelera)
    if hosts_a_ocultar:
        print(f"\n{Colors.ERROR}[!] Enviando {len(hosts_a_ocultar)} hosts únicos a la papelera. Por favor espere...{Colors.RESET}")
        hide_hosts(list(hosts_a_ocultar), falcon)
        print(f"{Colors.SUCCESS}[+] Proceso de ocultamiento finalizado.{Colors.RESET}")
    else:
        print(f"\n{Colors.SUCCESS}[-] No se encontraron duplicados para ocultar.{Colors.RESET}")

    # Exportar resultados a Excel
    excel_path = ruta_carpeta / "Depurados_Reporte.xlsx"
    with pd.ExcelWriter(excel_path) as writer:
        if not dep_name.empty:
            dep_name.to_excel(writer, sheet_name="dup_name", index=False)
        if not dep_serial.empty:
            dep_serial.to_excel(writer, sheet_name="dup_serial", index=False)
        if dep_name.empty and dep_serial.empty:
            pd.DataFrame({'Mensaje': ['Sin duplicados']}).to_excel(writer, sheet_name="Sin_Duplicados", index=False)
            
    print(f"\n{Colors.RESULT}[+] Reporte final generado en: {excel_path}{Colors.RESET}")

# =========================================================================
# LÓGICA PRINCIPAL (ORQUESTADOR)
# =========================================================================
def main() -> None:
    while True:
        opcion = menu_principal()

        if opcion == "0":
            print(f"\n{Colors.GREEN}Limpiando memoria y cerrando... ¡Hasta pronto!{Colors.RESET}")
            os.environ.pop("FALCON_CLIENT_ID", None)
            os.environ.pop("FALCON_CLIENT_SECRET", None)
            sys.exit(0)

        elif opcion == "1":
            client_id, client_secret = pedir_credenciales()
            run_tags_module(client_id, client_secret)
            input(f"\n{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")

        elif opcion == "2":
            client_id, client_secret = pedir_credenciales()
            run_rtr_module(client_id, client_secret)

        elif opcion == "3":
            client_id, client_secret = pedir_credenciales()
            run_ioc_module(client_id, client_secret)
            input(f"\n{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")

        elif opcion == "4":
            client_id, client_secret = pedir_credenciales()
            run_rtr_zip_sandbox_module(client_id, client_secret)
            input(f"\n{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")

        elif opcion == "5":
            client_id, client_secret = pedir_credenciales()
            run_rtr_universal_module(client_id=client_id, client_secret=client_secret)

        elif opcion == "6":
            client_id, client_secret = pedir_credenciales()
            run_upload_sensors_module(client_id, client_secret)

        elif opcion == "7":
            client_id, client_secret = pedir_credenciales()
            run_virustotal_hash_module(client_id, client_secret)
            input(f"\n{Colors.CYAN}Presione Enter para volver...{Colors.RESET}")

        elif opcion == "8":
            client_id, client_secret = pedir_credenciales()
            run_delete_hash_module(client_id, client_secret)
            
        # ==========================================
        # NUEVO MÓDULO OPERATIVO (9)
        # ==========================================
        elif opcion == "9":
            client_id, client_secret = pedir_credenciales() 
            run_depurador_module(client_id, client_secret)
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")

        # ==========================================
        # MÓDULOS DE REPORTES (Subprocesos)
        # ==========================================
        elif opcion == "10":
            pedir_credenciales() 
            print(f"\n{Colors.CYAN}Iniciando módulo de Riesgo de Identidad...{Colors.RESET}")
            subprocess.run(["python3", "Identity_Risk_v2.py"])
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")
            
        elif opcion == "11":
            pedir_credenciales()
            print(f"\n{Colors.CYAN}Iniciando módulo de Auditoría de Prevención...{Colors.RESET}")
            subprocess.run(["pwsh", "-File", "Politicas_Actuales_v2.ps1"])
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")
            
        elif opcion == "12":
            pedir_credenciales()
            print(f"\n{Colors.CYAN}Iniciando módulo de Exclusiones de Device Control...{Colors.RESET}")
            subprocess.run(["pwsh", "-File", "Report_CS_DeviceControl.ps1"])
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")

        elif opcion == "13":
            pedir_credenciales()
            print(f"\n{Colors.CYAN}Iniciando módulo de Exclusiones Globales (ML/SV/IOA)...{Colors.RESET}")
            subprocess.run(["pwsh", "-File", "Reporte_Exclusiones_Globales.ps1"])
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")

        elif opcion == "14":
            pedir_credenciales()
            print(f"\n{Colors.CYAN}Iniciando módulo de Accesos de Usuarios...{Colors.RESET}")
            subprocess.run(["pwsh", "-File", "Reporte_Usuarios_Consola.ps1"])
            input(f"\n{Colors.CYAN}Presione Enter para volver al menú...{Colors.RESET}")

        else:
            print(f"{Colors.ERROR}Opción inválida.{Colors.RESET}")
            time.sleep(1)

if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print(f"\n{Colors.WARNING}Interrumpido por el usuario{Colors.RESET}")
        sys.exit(0)